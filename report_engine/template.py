"""Template loading (.pptx / .potx), slide clearing, and the generated 16:9 brand template fallback."""
from __future__ import annotations

import io
import zipfile
from pathlib import Path

from pptx import Presentation

from .config import load_config
from .paths import design_dir, resolve_path

_TEMPLATE_CT = b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
_PRES_CT = b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
TEMPLATE_EXT = (".pptx", ".potx")
BRAND_NAME = "_brand_16x9.pptx"


def _potx_as_pptx(path: Path) -> io.BytesIO:
    buf = io.BytesIO()
    with zipfile.ZipFile(str(path)) as src, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(_TEMPLATE_CT, _PRES_CT)
            dst.writestr(item, data)
    buf.seek(0)
    return buf


def open_presentation(path):
    path = Path(path)
    if path.suffix.lower() == ".potx":
        return Presentation(_potx_as_pptx(path))
    return Presentation(str(path))


def remove_all_slides(prs) -> None:
    """Company templates often ship with sample slides; start from an empty deck."""
    sld_lst = prs.slides._sldIdLst
    for sld in list(sld_lst):
        r_id = sld.rId
        sld_lst.remove(sld)
        prs.part.drop_rel(r_id)


def ensure_brand_template(cfg=None) -> Path:
    """The generated 16:9 brand template; rebuilt whenever theme.yaml (or the logo) changes."""
    from .brand_template import BRAND_VERSION, build_brand_template
    from .theme import load_theme

    theme = load_theme(cfg)
    dest = design_dir() / "templates" / BRAND_NAME
    stamp = dest.with_suffix(".hash")
    fp = "%s-%s" % (BRAND_VERSION, theme.fingerprint())
    if not dest.exists() or not stamp.exists() or stamp.read_text(encoding="utf-8").strip() != fp:
        build_brand_template(dest, theme)
        stamp.write_text(fp, encoding="utf-8")
    return dest


def resolve_template_path(explicit=None, cfg=None) -> Path:
    cfg = cfg or load_config()
    tdir = design_dir() / "templates"
    if explicit:
        p = resolve_path(explicit, base=design_dir())
        if not p.exists():
            raise FileNotFoundError("找不到範本：%s" % explicit)
        return p
    if cfg.get("template"):
        p = resolve_path(cfg["template"], base=design_dir())
        if p.exists():
            return p
    if tdir.exists():
        cands = sorted(
            x for x in tdir.iterdir()
            if x.suffix.lower() in TEMPLATE_EXT and not x.name.startswith(("_builtin", "_brand", "~$"))
        )
        if cands:
            return cands[0]
    return ensure_brand_template(cfg)
