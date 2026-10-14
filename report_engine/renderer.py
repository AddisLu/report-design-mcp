"""Render normalized slides onto the template: placeholders first, grid regions for sketched layouts."""
from __future__ import annotations

import math
import re
import unicodedata

from lxml import etree
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_MARK
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from .template import open_presentation, remove_all_slides
from .template_inspector import TEXT_KINDS, TITLE_TYPES, inspect_presentation
from .theme import load_theme, mix

EMU_PT = 12700
EMU_IN = 914400
GUTTER = int(EMU_IN * 0.08)
RICH_RE = re.compile(r"(\*\*.+?\*\*|!!.+?!!)")
_NUMERIC_CELL = re.compile(r"^[+\-]?[\d,]+(\.\d+)?\s*%?$")
_BU_SUCCESSORS = (qn("a:tabLst"), qn("a:defRPr"), qn("a:extLst"))
_BENCHMARK_RE = re.compile(u"目標|計畫|計劃|預算|預估|基準|去年|上期|target|plan|budget|baseline|goal", re.I)
_UP_RE = re.compile(u"^\\s*(▲|↑|\\+|增加|上升|成長|提升|提高)")
_DOWN_RE = re.compile(u"^\\s*(▼|↓|-|−|減少|下降|衰退|降低)")
_NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
_SLIDENUM_ID = "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}"

CHART_MAP = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "stacked": XL_CHART_TYPE.COLUMN_STACKED,
    "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "area": XL_CHART_TYPE.AREA,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}


def plain_text(s) -> str:
    return RICH_RE.sub(lambda m: m.group(0)[2:-2], s or "")


def em_width(s: str) -> float:
    w = 0.0
    for ch in s:
        if unicodedata.east_asian_width(ch) in ("W", "F"):
            w += 1.0
        elif ch == " ":
            w += 0.28
        elif ch.isupper() or ch.isdigit():
            w += 0.6
        else:
            w += 0.5
    return w


def _text_height(items, width_pt: float) -> float:
    """items: (text, size_pt, indent_pt). Rough line-wrap estimate, deliberately a bit pessimistic."""
    total = 0.0
    for text, size, indent in items:
        avail = max(width_pt - indent, size * 3)
        lines = 0
        for seg in (text or " ").split("\n"):
            lines += max(1, int(math.ceil(em_width(seg) * size / avail)))
        total += lines * size * 1.2 + size * 0.3
    return total


def _remove(shape) -> None:
    el = shape._element
    el.getparent().remove(el)


def _lst_size(el, lvl):
    if el is None:
        return None
    d = el.find(".//%s/%s/%s" % (qn("a:lstStyle"), qn("a:lvl%dpPr" % lvl), qn("a:defRPr")))
    if d is not None and d.get("sz"):
        return int(d.get("sz")) / 100.0
    return None


def _master_sizes(master):
    tx = master._element.find(qn("p:txStyles"))

    def get(style, lvl, default):
        if tx is None:
            return default
        st = tx.find(qn(style))
        if st is None:
            return default
        d = st.find("%s/%s" % (qn("a:lvl%dpPr" % lvl), qn("a:defRPr")))
        return int(d.get("sz")) / 100.0 if d is not None and d.get("sz") else default

    title = get("p:titleStyle", 1, 36.0)
    body = [get("p:bodyStyle", i, d) for i, d in ((1, 24.0), (2, 20.0), (3, 18.0))]
    return title, body


def _insert_bu(pPr, el) -> None:
    for i, child in enumerate(pPr):
        if child.tag in _BU_SUCCESSORS:
            pPr.insert(i, el)
            return
    pPr.append(el)


def _clear_bu(pPr) -> None:
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buBlip", "a:buFont", "a:buClr"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)


def _no_bullet(para) -> None:
    pPr = para._p.get_or_add_pPr()
    pPr.set("marL", "0")
    pPr.set("indent", "0")
    _clear_bu(pPr)
    _insert_bu(pPr, etree.Element(qn("a:buNone")))


def _set_bullet(para, level: int, size: float) -> None:
    pPr = para._p.get_or_add_pPr()
    ind = int(size * EMU_PT * 1.1)
    pPr.set("marL", str(ind * (level + 1)))
    pPr.set("indent", str(-ind))
    _clear_bu(pPr)
    bu_clr = etree.SubElement(etree.Element(qn("a:buClr")), qn("a:schemeClr"))
    bu_clr.set("val", "accent1")
    bu_font = etree.Element(qn("a:buFont"))
    bu_font.set("typeface", "Arial")
    bu_char = etree.Element(qn("a:buChar"))
    bu_char.set("char", u"•" if level != 1 else u"–")
    _insert_bu(pPr, bu_clr.getparent())
    _insert_bu(pPr, bu_font)
    _insert_bu(pPr, bu_char)


def _set_color(color_format, color) -> None:
    """color: 6-digit hex string or an MSO_THEME_COLOR member."""
    if isinstance(color, str):
        color_format.rgb = RGBColor.from_string(color)
    else:
        color_format.theme_color = color


def _plain_shape(shp, color=MSO_THEME_COLOR.ACCENT_1) -> None:
    style = shp._element.find(qn("p:style"))
    if style is not None:
        shp._element.remove(style)
    shp.fill.solid()
    _set_color(shp.fill.fore_color, color)
    shp.line.fill.background()
    shp.shadow.inherit = False


def _muted(font) -> None:
    font.color.theme_color = MSO_THEME_COLOR.TEXT_1
    font.color.brightness = 0.4


def _direction(text):
    t = plain_text(text or "")
    if _UP_RE.match(t):
        return "up"
    if _DOWN_RE.match(t):
        return "down"
    return None


def _no_frame(el, tag):
    """Give a chart element (c:chartSpace / c:plotArea) an explicit no-fill, no-line shape style."""
    spPr = el.find(qn("c:spPr"))
    if spPr is None:
        spPr = etree.SubElement(el, qn("c:spPr"))
        anchor = el.find(qn(tag))
        if anchor is not None:
            anchor.addnext(spPr)
    for child in list(spPr):
        spPr.remove(child)
    etree.SubElement(spPr, qn("a:noFill"))
    etree.SubElement(etree.SubElement(spPr, qn("a:ln")), qn("a:noFill"))


class DeckRenderer:
    def __init__(self, template_path, cfg):
        self.cfg = cfg
        self.prs = open_presentation(template_path)
        remove_all_slides(self.prs)
        self.info = inspect_presentation(self.prs)
        self.by_key = {L["key"]: L for L in self.info["layouts"]}
        self.sw, self.sh = int(self.prs.slide_width), int(self.prs.slide_height)
        self.min_pt = float(cfg.get("min_font_pt", 12))
        self.theme = load_theme(cfg)
        self.alert_rgb = RGBColor.from_string(self.theme.alert)
        self.gap = self.theme.inch("gap") or GUTTER
        self.meta = {}
        self.warnings = []
        self._where = ""
        master = self.prs.slide_masters[0]
        self.title_pt, self.body_pts = _master_sizes(master)
        self._calc_area(master)

    # -- geometry -------------------------------------------------------
    def _calc_area(self, master) -> None:
        mx = self.theme.inch("margin_x") or int(self.sw * 0.05)
        mt = self.theme.inch("margin_top") or int(self.sh * 0.06)
        bottom = self.sh - mt
        for ph in master.placeholders:
            t = ph.placeholder_format.type
            if ph.left is None or ph.top is None:
                continue
            if t in TITLE_TYPES:
                if 0 < ph.left < self.sw * 0.2:
                    mx = int(ph.left)
                if 0 < ph.top < self.sh * 0.2:
                    mt = int(ph.top)
            if t in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE) and ph.top > self.sh * 0.75:
                bottom = min(bottom, int(ph.top) - self.gap // 2)
        self.area = (mx, mt, self.sw - 2 * mx, bottom - mt)

    def grid_rect(self, g):
        ax, ay, aw, ah = self.area
        x, y, w, h = g
        g = self.gap
        # outer edges of the region grid sit flush with the content margins; gaps only between regions
        x0 = ax + x / 12.0 * aw + (g / 2.0 if x > 0.01 else 0)
        y0 = ay + y / 8.0 * ah + (g / 2.0 if y > 0.01 else 0)
        x1 = ax + (x + w) / 12.0 * aw - (g / 2.0 if x + w < 11.99 else 0)
        y1 = ay + (y + h) / 8.0 * ah - (g / 2.0 if y + h < 7.99 else 0)
        return int(x0), int(y0), max(int(x1 - x0), EMU_IN // 4), max(int(y1 - y0), EMU_IN // 4)

    def _title_bottom_grid(self) -> float:
        L = self.by_key[self.info["with_title"]]
        if not L.get("title_bottom"):
            return 0.0
        ay, ah = self.area[1], self.area[3]
        return min(3.0, max(0.0, (L["title_bottom"] - ay) / float(ah) * 8 + 0.15))

    # -- entry ----------------------------------------------------------
    def render_all(self, nslides):
        out = []
        for ns in nslides:
            start = len(self.prs.slides)
            self._where = ns.get("where", "")
            try:
                if ns["mode"] == "custom":
                    self._custom(ns, ns["regions"])
                elif ns["mode"] == "section":
                    self._section(ns)
                else:
                    self._layout_slide(ns)
            except Exception as e:  # keep the rest of the deck
                self.warnings.append("%s：繪製失敗（%s: %s）" % (self._where, type(e).__name__, e))
            out.append(list(range(start + 1, len(self.prs.slides) + 1)))
        return out

    def save(self, path) -> None:
        self.prs.save(str(path))

    # -- layout (placeholder) slides -----------------------------------
    @staticmethod
    def _headline_only(blocks) -> bool:
        return (len(blocks) == 1 and blocks[0]["type"] == "text"
                and all(p.get("headline") for p in blocks[0]["paras"]))

    def _choose(self, ns):
        blocks = ns["blocks"]
        key = ns.get("layout")
        if key:
            L = self.by_key.get(key)
            if L is None:
                for k, v in self.by_key.items():
                    if k.lower() == str(key).lower():
                        L = v
            if L is not None:
                return L
            self.warnings.append("%s：範本沒有版型「%s」，已自動挑選" % (self._where, key))
        cover = self.info.get("cover")
        if cover and (not blocks or self._headline_only(blocks)) and ns["idx"] == 1 and not ns.get("label"):
            return self.by_key[cover]
        k = self.info["by_count"].get(max(1, len(blocks)))
        return self.by_key[k] if k else None

    def _layout_slide(self, ns) -> None:
        L = self._choose(ns)
        blocks = list(ns["blocks"])
        cover_fit = L is not None and L.get("subtitle_idx") is not None and self._headline_only(blocks)
        if L is None or (len(blocks) > len(L["content_idx"]) and not cover_fit):
            if L is not None:
                self.warnings.append("%s：內容區塊（%d）多於版型「%s」的內容區（%d），已自動改用自訂排版"
                                     % (self._where, len(blocks), L["key"], len(L["content_idx"])))
            else:
                self.warnings.append("%s：範本沒有剛好 %d 個內容區的版型，已自動改用自訂排版" % (self._where, len(blocks)))
            self._custom(ns, self._auto_regions(blocks, bool(ns.get("title"))))
            return
        slide = self.prs.slides.add_slide(L["obj"])
        phs = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
        title_ph = phs.get(L["title_idx"]) if L["title_idx"] is not None else None
        if title_ph is not None:
            if ns.get("title"):
                self._set_title(title_ph, ns["title"], L["obj"])
            else:
                _remove(title_ph)
        elif ns.get("title"):
            self.warnings.append("%s：版型「%s」沒有標題區，標題未放入" % (self._where, L["key"]))
        sub_ph = phs.get(L["subtitle_idx"]) if L.get("subtitle_idx") is not None else None
        if sub_ph is not None:
            if cover_fit:
                self._fill_text_ph(sub_ph, blocks[0]["paras"], L["obj"], allow_split=False, plain_headline=True)
                self._cover_meta(sub_ph)
                blocks = []
            elif L["key"] == self.info.get("cover") and not blocks:
                self._cover_meta(sub_ph, first=True)
            else:
                _remove(sub_ph)
        content_phs = [phs[i] for i in L["content_idx"] if i in phs]
        kinds = [c["kind"] for c in L["content"]]
        leftover = None
        for i, (ph, block) in enumerate(zip(content_phs, blocks)):
            if block["type"] == "text" and kinds[i] in TEXT_KINDS:
                leftover = self._fill_text_ph(ph, block["paras"], L["obj"], allow_split=True)
            else:
                rect = (int(ph.left), int(ph.top), int(ph.width), int(ph.height))
                _remove(ph)
                self._draw(slide, rect, block)
        for ph in content_phs[len(blocks):]:
            _remove(ph)
        self._notes(slide, ns.get("notes"))
        if L["key"] != self.info.get("cover"):
            self._chrome(slide, L["obj"], ns)
        guard = 0
        while leftover and guard < 6:
            guard += 1
            leftover = self._continuation(ns, L, leftover)

    def _continuation(self, ns, L, paras):
        slide = self.prs.slides.add_slide(L["obj"])
        phs = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
        title_ph = phs.get(L["title_idx"]) if L["title_idx"] is not None else None
        if title_ph is not None:
            self._set_title(title_ph, (ns.get("title") or "") + u"（續）", L["obj"])
        if L.get("subtitle_idx") in phs:
            _remove(phs[L["subtitle_idx"]])
        content_phs = [phs[i] for i in L["content_idx"] if i in phs]
        if not content_phs:
            return None
        left = self._fill_text_ph(content_phs[0], paras, L["obj"], allow_split=True)
        for ph in content_phs[1:]:
            _remove(ph)
        self._chrome(slide, L["obj"], ns)
        return left

    def _ph_sizes(self, ph, layout, default):
        idx = ph.placeholder_format.idx
        sizes = list(default)
        lph = None
        for p in layout.placeholders:
            if p.placeholder_format.idx == idx:
                lph = p
                break
        for el in ((lph._element if lph is not None else None), ph._element):
            for lvl in (1, 2, 3):
                sz = _lst_size(el, lvl)
                if sz:
                    sizes[lvl - 1] = sz
        for i in (1, 2):
            sizes[i] = min(sizes[i], sizes[i - 1])
        return sizes

    def _fit(self, height_fn, h_pt, base0, min_pt=None):
        """Largest scale that fits; candidate sizes come from the type scale so the deck stays consistent."""
        floor = min(float(base0), float(min_pt or self.min_pt))
        cands = [1.0] + [st / float(base0) for st in reversed(self.theme.steps) if floor - 1e-6 <= st < base0 - 1e-6]
        for s in cands:
            if height_fn(s) <= h_pt:
                return s
        return None

    def _set_title(self, ph, text, layout) -> None:
        base = self._ph_sizes(ph, layout, [self.title_pt] * 3)[0]
        w_pt = ph.width / float(EMU_PT) - 14.4
        h_pt = ph.height / float(EMU_PT) - 7.2
        plain = plain_text(text)

        def height(s):
            size = base * s
            lines = max(1, int(math.ceil(em_width(plain) * size / max(w_pt, size))))
            return lines * size * 1.1

        scale = self._fit(height, h_pt, base, min_pt=18)
        if scale is None:
            scale = min(1.0, 18.0 / base)
            self.warnings.append("%s：標題太長，建議精簡到 30 字內" % self._where)
        tf = ph.text_frame
        tf.clear()
        self._runs(tf.paragraphs[0], text, base * scale if scale < 0.999 else None)

    def _fill_text_ph(self, ph, paras, layout, allow_split, plain_headline=False):
        base = self._ph_sizes(ph, layout, self.body_pts)
        w_pt = ph.width / float(EMU_PT) - 14.4
        h_pt = (ph.height / float(EMU_PT) - 7.2) * 0.95

        def items(ps, s):
            return [(plain_text(p["text"]), base[min(p["level"], 2)] * s,
                     (p["level"] + 1) * 20.0 if p.get("bullet", True) and not p.get("headline") else 0.0)
                    for p in ps]

        scale = self._fit(lambda s: _text_height(items(paras, s), w_pt), h_pt, base[0])
        leftover = None
        if scale is None:
            min_s = min(1.0, self.min_pt / base[0])
            if allow_split and len(paras) > 1:
                k = len(paras)
                while k > 1 and _text_height(items(paras[:k], min_s), w_pt) > h_pt:
                    k -= 1
                leftover, paras = paras[k:], paras[:k]
                scale = self._fit(lambda s: _text_height(items(paras, s), w_pt), h_pt, base[0]) or min_s
            else:
                scale = min_s
                self.warnings.append("%s：文字過多，已縮到最小字級，仍可能超出版面" % self._where)
        explicit = scale < 0.999
        tf = ph.text_frame
        tf.clear()
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            lvl = min(p["level"], 2)
            if lvl:
                para.level = lvl
            if p.get("headline") or not p.get("bullet", True):
                _no_bullet(para)
            size = base[lvl] * scale if explicit else None
            self._runs(para, p["text"], size, bold=bool(p.get("headline")) and not plain_headline, emph=p.get("emph"))
        return leftover

    # -- custom (grid) slides ------------------------------------------
    def _auto_regions(self, blocks, has_title):
        top = self._title_bottom_grid() if has_title else 0.0
        H = 8.0 - top
        regs = []
        kp = [b for b in blocks if b["type"] == "kpis"]
        rest = [b for b in blocks if b["type"] != "kpis"]
        if kp and rest:
            band = min(2.0, H * 0.3)
            w = 12.0 / len(kp)
            for i, b in enumerate(kp):
                regs.append(((i * w, top, w, band), b))
            top += band
            H -= band
        elif kp:
            rest = kp
        n = len(rest)
        if n == 1:
            regs.append(((0, top, 12, H), rest[0]))
        elif n == 2:
            a = 5 if rest[0]["type"] == "text" else 6
            regs += [((0, top, a, H), rest[0]), ((a, top, 12 - a, H), rest[1])]
        elif n == 3:
            regs += [((0, top, 4, H), rest[0]), ((4, top, 8, H / 2), rest[1]), ((4, top + H / 2, 8, H / 2), rest[2])]
        elif n > 3:
            rows = int(math.ceil(n / 2.0))
            for i, b in enumerate(rest):
                regs.append((((i % 2) * 6, top + (i // 2) * H / rows, 6, H / rows), b))
        return [{"rect": r, "block": b} for r, b in regs]

    def _custom(self, ns, regions) -> None:
        has_title = bool(ns.get("title"))
        L = self.by_key[self.info["with_title"] if has_title else self.info["blank"]]
        slide = self.prs.slides.add_slide(L["obj"])
        title_bottom = self.area[1]
        for ph in list(slide.placeholders):
            if has_title and ph.placeholder_format.type in TITLE_TYPES and title_bottom == self.area[1]:
                self._set_title(ph, ns["title"], L["obj"])
                title_bottom = int(ph.top + ph.height)
            else:
                _remove(ph)
        for reg in regions:
            x, y, w, h = self.grid_rect(reg["rect"])
            if has_title and y < title_bottom + self.gap // 2:
                h = max(h - (title_bottom + self.gap - y), EMU_IN // 2)
                y = title_bottom + self.gap
            self._draw(slide, (x, y, w, h), reg["block"])
        self._notes(slide, ns.get("notes"))
        self._chrome(slide, L["obj"], ns)

    # -- cover / section / page chrome ------------------------------------
    def _cover_meta(self, sub_ph, first=False) -> None:
        line = u"　·　".join(str(v) for v in (self.meta.get("author"), self.meta.get("date")) if v)
        if not line:
            if first:
                _remove(sub_ph)
            return
        tf = sub_ph.text_frame
        if first:
            tf.clear()
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        _no_bullet(para)
        para.space_before = Pt(18)
        self._runs(para, line, self.theme.size("caption") + 2)

    def _section(self, ns) -> None:
        key = self.info.get("section") or self.info.get("cover")
        if not key:
            blocks = [{"type": "headline", "text": ns.get("subtitle") or ""}] if ns.get("subtitle") else []
            self._custom(ns, self._auto_regions(blocks, True))
            return
        L = self.by_key[key]
        slide = self.prs.slides.add_slide(L["obj"])
        phs = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
        title_ph = phs.get(L["title_idx"]) if L["title_idx"] is not None else None
        if title_ph is not None:
            self._set_title(title_ph, ns.get("title") or "", L["obj"])
        sub_idx = L.get("subtitle_idx")
        if sub_idx is None and L["content_idx"]:
            sub_idx = L["content_idx"][0]  # company "Section Header" layouts often use a body placeholder
        sub_ph = phs.get(sub_idx) if sub_idx is not None else None
        keep = {L["title_idx"], sub_idx}
        for ph in list(slide.placeholders):
            if ph.placeholder_format.idx not in keep:
                _remove(ph)
        if sub_ph is not None:
            if ns.get("subtitle"):
                self._fill_text_ph(sub_ph, [{"text": ns["subtitle"], "level": 0, "bullet": False, "headline": True}],
                                   L["obj"], allow_split=False, plain_headline=True)
            else:
                _remove(sub_ph)
        self._notes(slide, ns.get("notes"))

    def _chrome(self, slide, layout, ns) -> None:
        """Footer line (資料來源 or footer text) + slide number, using the layout's own placeholders when it has them."""
        src = ns.get("source")
        text = (u"資料來源：" + plain_text(src)) if src else self.theme.footer_text
        lphs = {}
        for ph in layout.placeholders:
            lphs.setdefault(ph.placeholder_format.type, ph)
        cap = self.theme.size("caption")
        num = len(self.prs.slides)
        y = self.sh - int(EMU_IN * 0.48)
        if text:
            p = self._chrome_para(slide, lphs.get(PP_PLACEHOLDER.FOOTER),
                                  (self.area[0], y, int(self.area[2] * 0.7), int(EMU_IN * 0.3)), PP_ALIGN.LEFT)
            self._runs(p, text, cap, color=self.theme.neutral(500))
        if self.theme.page_number:
            p = self._chrome_para(slide, lphs.get(PP_PLACEHOLDER.SLIDE_NUMBER),
                                  (self.area[0] + self.area[2] - EMU_IN, y, EMU_IN, int(EMU_IN * 0.3)), PP_ALIGN.RIGHT)
            fld = etree.SubElement(p._p, qn("a:fld"), id=_SLIDENUM_ID, type="slidenum")
            rpr = etree.SubElement(fld, qn("a:rPr"), lang="zh-TW", sz=str(int(cap * 100)))
            etree.SubElement(etree.SubElement(rpr, qn("a:solidFill")), qn("a:srgbClr")).set("val", self.theme.neutral(500))
            etree.SubElement(fld, qn("a:t")).text = str(num)

    def _chrome_para(self, slide, layout_ph, rect, align):
        if layout_ph is not None:
            idx = layout_ph.placeholder_format.idx
            slide.shapes.clone_placeholder(layout_ph)
            ph = next(p for p in slide.placeholders if p.placeholder_format.idx == idx)
            tf = ph.text_frame
        else:
            _tb, tf = self._textbox(slide, rect, MSO_ANCHOR.MIDDLE)
        tf.clear()
        para = tf.paragraphs[0]
        para.alignment = align
        return para

    # -- blocks ---------------------------------------------------------
    def _draw(self, slide, rect, block) -> None:
        t = block["type"]
        th = self.theme
        if t == "text":
            self._draw_text(slide, rect, block["paras"], tuple(th.size("body", k) for k in range(3)))
        elif t == "headline":
            hs = th.size("headline")
            self._draw_text(slide, rect, [{"text": block["text"], "level": 0, "bullet": False, "headline": True}],
                            (hs, hs, hs), anchor=MSO_ANCHOR.MIDDLE, min_pt=14)
        elif t == "note":
            ns = th.size("note")
            self._draw_text(slide, rect, [{"text": block["text"], "level": 0, "bullet": False}],
                            (ns, ns, ns), min_pt=9, muted=True)
        elif t == "kpis":
            self._draw_kpis(slide, rect, block["items"])
        elif t == "status":
            self._draw_status(slide, rect, block["items"])
        elif t == "chart":
            self._draw_chart(slide, rect, block)
        elif t == "table":
            self._draw_table(slide, rect, block)
        elif t == "image":
            self._draw_image(slide, rect, block)

    def _textbox(self, slide, rect, anchor=MSO_ANCHOR.TOP):
        x, y, w, h = rect
        tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = 0  # text edges line up with the title and the grid
        tf.margin_top = tf.margin_bottom = Emu(45720)
        tf.vertical_anchor = anchor
        return tb, tf

    def _draw_text(self, slide, rect, paras, base, anchor=MSO_ANCHOR.TOP, min_pt=None, muted=False) -> None:
        w_pt = rect[2] / float(EMU_PT) - 2
        h_pt = (rect[3] / float(EMU_PT) - 7.2) * 0.95

        def height(s):
            return _text_height([(plain_text(p["text"]), base[min(p["level"], 2)] * s,
                                  (p["level"] + 1) * base[min(p["level"], 2)] * s * 1.1 if p.get("bullet", True) else 0.0)
                                 for p in paras], w_pt)

        scale = self._fit(height, h_pt, base[0], min_pt)
        if scale is None:
            scale = min(1.0, (min_pt or self.min_pt) / base[0])
            self.warnings.append("%s：區塊文字過多，已縮到最小字級，仍可能超出框線" % self._where)
        _tb, tf = self._textbox(slide, rect, anchor)
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            lvl = min(p["level"], 2)
            size = base[lvl] * scale
            if i:
                para.space_before = Pt(round(size * 0.3, 1))
            if p.get("bullet", True) and not p.get("headline"):
                _set_bullet(para, lvl, size)
            self._runs(para, p["text"], size, bold=bool(p.get("headline")), emph=p.get("emph"),
                       color=_muted if muted else None)

    def _runs(self, para, text, size=None, bold=False, emph=None, color=None) -> None:
        """color: callable(font), hex string, or theme colour; emphasis wins over it."""
        for part in RICH_RE.split(text or ""):
            if not part:
                continue
            e = emph
            if part.startswith("**") and part.endswith("**") and len(part) >= 4:
                part, e = part[2:-2], (emph or "accent")
            elif part.startswith("!!") and part.endswith("!!") and len(part) >= 4:
                part, e = part[2:-2], "alert"
            r = para.add_run()
            r.text = part.replace("\n", " ")
            f = r.font
            if size:
                f.size = Pt(self.theme.snap(size))
            if bold or e:
                f.bold = True
            if e in ("accent", "accent1"):
                f.color.theme_color = MSO_THEME_COLOR.ACCENT_1
            elif e == "accent2":
                f.color.theme_color = MSO_THEME_COLOR.ACCENT_2
            elif e in ("alert", "red"):
                f.color.rgb = self.alert_rgb
            elif e in ("good", "warn", "bad"):
                f.color.rgb = RGBColor.from_string(self.theme.status_text(e))
            elif callable(color):
                color(f)
            elif color is not None:
                _set_color(f.color, color)

    def _draw_kpis(self, slide, rect, items) -> None:
        n = len(items)
        if not n:
            return
        x, y, w, h = rect
        g = self.gap
        best = None
        for cols in range(1, min(n, 4) + 1):
            rows = int(math.ceil(n / float(cols)))
            score = abs(math.log(((w / float(cols)) / (h / float(rows))) / 2.0))
            if best is None or score < best[0]:
                best = (score, cols, rows)
        _s, cols, rows = best
        cw, ch = (w - g * (cols - 1)) // cols, (h - g * (rows - 1)) // rows
        pad = int(min(EMU_IN * 0.16, ch * 0.1))
        tick = ch > EMU_IN * 1.6
        w_pt = (cw - 2 * pad) / float(EMU_PT)
        h_pt = (ch - 2 * pad - (EMU_IN * 0.1 if tick else 0)) / float(EMU_PT)
        lsize = self.theme.snap(min(self.theme.size("kpi_label"), max(12.0, h_pt * 0.16)))
        lines = (1 if any(it.get("label") for it in items) else 0) + (1 if any(it.get("delta") for it in items) else 0)
        avail = h_pt - lines * lsize * 1.35
        # one value size for the whole row of cards, so the numbers line up
        vsize = self.theme.size("kpi_value")
        for it in items:
            val = plain_text(it["value"])
            vsize = min(vsize, avail / 1.2, w_pt * 0.95 / max(em_width(val), 1.0))
        vsize = self.theme.snap(max(vsize, 16.0))
        for i, it in enumerate(items):
            cx = x + (i % cols) * (cw + g)
            cy = y + (i // cols) * (ch + g)
            self._kpi_card(slide, (cx, cy, cw, ch), it, vsize, lsize, pad, tick)

    def _kpi_card(self, slide, rect, it, vsize, lsize, pad, tick) -> None:
        x, y, w, h = rect
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
        card.adjustments[0] = min(0.5, EMU_IN * 0.08 / float(min(w, h)))
        _plain_shape(card, self.theme.neutral(50))
        status = it.get("status")
        tick_color = self.theme.status(status) if status else MSO_THEME_COLOR.ACCENT_1
        if tick:
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x + pad), Emu(y + pad),
                                         Emu(int(EMU_IN * 0.36)), Emu(int(EMU_IN * 0.05)))
            _plain_shape(bar, tick_color)
            top = int(EMU_IN * 0.1)
        else:
            top = 0
        _tb, tf = self._textbox(slide, (x + pad, y + pad + top, w - 2 * pad, h - 2 * pad - top), MSO_ANCHOR.MIDDLE)
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        if it.get("label"):
            self._runs(tf.paragraphs[0], it["label"], lsize, color=self.theme.neutral(700))
            vp = tf.add_paragraph()
        else:
            vp = tf.paragraphs[0]
        vp.space_before = Pt(2)
        emph = it.get("emph") or (status if status in ("bad", "warn") else None)
        self._runs(vp, it["value"], vsize, bold=True, emph=emph, color=self.theme.neutral(900))
        if it.get("delta"):
            d = plain_text(it["delta"])
            direction = _direction(d)
            if direction:
                good = (it.get("good") or "up") == direction
                color = self.theme.status("good" if good else "bad")
                if d.lstrip()[:1] not in (u"▲", u"▼", u"↑", u"↓"):
                    d = (u"▲ " if direction == "up" else u"▼ ") + d.strip()
            else:
                color = self.theme.neutral(500)
            dp = tf.add_paragraph()
            dp.space_before = Pt(2)
            self._runs(dp, d, lsize, color=color)

    def _draw_status(self, slide, rect, items) -> None:
        if not items:
            return
        x, y, w, h = rect
        n = len(items)
        row_h = min(h / float(n), EMU_IN * 0.8)
        if row_h < EMU_IN * 0.32:
            keep = max(1, int(h // (EMU_IN * 0.32)))
            self.warnings.append(u"%s：狀態清單 %d 項放不下，只顯示前 %d 項（可拆頁）" % (self._where, n, keep))
            items, n, row_h = items[:keep], keep, h / float(keep)
        rh_pt = row_h / float(EMU_PT)
        size = self.theme.snap(max(11.0, min(self.theme.size("body", 1), rh_pt * 0.36)))
        small = self.theme.snap(max(10.0, size - 2))
        has_owner = any(it.get("owner") for it in items)
        has_note = any(it.get("note") for it in items)
        pill_w = int(EMU_IN * 1.15)
        owner_w = int(w * 0.16) if has_owner else 0
        name_w = int((w - pill_w - owner_w) * (0.36 if has_note else 1.0))
        note_w = w - pill_w - owner_w - name_w
        dot = int(size * 0.62 * EMU_PT)
        for i, it in enumerate(items):
            ry = int(y + i * row_h)
            key = it.get("status")
            color = self.theme.status(key) if key else self.theme.neutral(300)
            oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(x), Emu(int(ry + (row_h - dot) / 2)), Emu(dot), Emu(dot))
            _plain_shape(oval, color)
            label = it.get("label") or {"good": u"正常", "warn": u"注意", "bad": u"延遲"}.get(key, u"—")
            cells = [((x + dot + int(EMU_IN * 0.08), pill_w - dot - int(EMU_IN * 0.08)), label, small, True,
                      self.theme.status_text(key) if key else self.theme.neutral(500), PP_ALIGN.LEFT),
                     ((x + pill_w, name_w), it.get("name") or "", size, True, self.theme.neutral(900), PP_ALIGN.LEFT)]
            if has_note:
                cells.append(((x + pill_w + name_w, note_w), it.get("note") or "", small, False,
                              self.theme.neutral(700), PP_ALIGN.LEFT))
            if has_owner:
                cells.append(((x + w - owner_w, owner_w), it.get("owner") or "", small, False,
                              self.theme.neutral(500), PP_ALIGN.RIGHT))
            for (cx, cw), text, sz, bold, color_, align in cells:
                _tb, tf = self._textbox(slide, (cx, ry, cw, int(row_h)), MSO_ANCHOR.MIDDLE)
                tf.margin_top = tf.margin_bottom = 0
                p = tf.paragraphs[0]
                p.alignment = align
                self._runs(p, text, sz, bold=bold, color=color_)
            if i < n - 1:
                ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(x), Emu(int(ry + row_h)),
                                                Emu(x + w), Emu(int(ry + row_h)))
                style = ln._element.find(qn("p:style"))
                if style is not None:
                    ln._element.remove(style)
                ln.line.color.rgb = RGBColor.from_string(self.theme.neutral(100))
                ln.line.width = Pt(0.75)

    def _chart_title(self, slide, rect, title):
        x, y, w, h = rect
        size = self.theme.size("body", 2)
        th = int(size * 1.7 * EMU_PT)
        _tb, tf = self._textbox(slide, (x, y, w, th), MSO_ANCHOR.TOP)
        tf.margin_left = tf.margin_right = 0
        self._runs(tf.paragraphs[0], title, size, bold=True, color=self.theme.neutral(700))
        return x, y + th, w, h - th

    def _draw_chart(self, slide, rect, b) -> None:
        ctype = b["chart_type"]
        series = b["series"]
        pie = ctype in ("pie", "doughnut")
        bars = ctype in ("column", "bar", "stacked", "stacked_bar")
        if pie and len(series) > 1:
            self.warnings.append("%s：圓餅圖只能有一組數列，已只用「%s」" % (self._where, series[0]["name"]))
            series = series[:1]
        if b.get("title"):
            rect = self._chart_title(slide, rect, b["title"])
        x, y, w, h = rect
        pct = all(s.get("percent") for s in series)
        data = CategoryChartData(number_format="0%" if pct else "General")
        data.categories = [str(c) for c in b["categories"]]
        for s in series:
            data.add_series(str(s["name"] or " "), [None if v is None else float(v) for v in s["values"]])
        chart = slide.shapes.add_chart(CHART_MAP[ctype], Emu(x), Emu(y), Emu(w), Emu(h), data).chart
        th = self.theme
        chart.font.size = Pt(th.size("chart"))
        chart.font.color.rgb = RGBColor.from_string(th.neutral(700))
        chart.has_title = False
        _no_frame(chart._chartSpace, "c:chart")
        _no_frame(chart._chartSpace.find(qn("c:chart")).find(qn("c:plotArea")), "c:valAx")
        multi = len(series) > 1 or pie
        chart.has_legend = multi
        if multi:
            chart.legend.position = XL_LEGEND_POSITION.TOP
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(th.size("chart"))

        # colours: benchmark series (目標/預算/去年…) in grey, the rest from the palette; highlight → primary
        palette = th.chart_colors
        highlight = set(str(c) for c in (b.get("highlight") or []))
        hl_color = th.alert if b.get("highlight_color") in ("alert", "red", "bad") else th.primary
        plot = chart.plots[0]
        k = 0
        first_main = None
        for i, s in enumerate(series):
            ser = plot.series[i]
            bench = len(series) > 1 and bool(_BENCHMARK_RE.search(str(s["name"])))
            color = th.neutral(300) if bench else palette[k % len(palette)]
            if not bench:
                k += 1
                first_main = i if first_main is None else first_main
            if pie:
                for j, cat in enumerate(b["categories"]):
                    pt = ser.points[j]
                    pt.format.fill.solid()
                    c = palette[j % len(palette)]
                    if highlight:
                        c = hl_color if str(cat) in highlight else th.neutral(300)
                    pt.format.fill.fore_color.rgb = RGBColor.from_string(c)
                    pt.format.line.color.rgb = RGBColor.from_string("FFFFFF")
                continue
            if ctype == "line":
                ser.smooth = False
                ser.format.line.color.rgb = RGBColor.from_string(color)
                ser.format.line.width = Pt(1.5 if bench else 2.5)
                if bench:
                    ser.format.line.dash_style = 4  # MSO_LINE.DASH
                ser.marker.style = XL_MARKER_STYLE.CIRCLE
                ser.marker.size = 6
                ser.marker.format.fill.solid()
                ser.marker.format.fill.fore_color.rgb = RGBColor.from_string(color)
                ser.marker.format.line.color.rgb = RGBColor.from_string(color)
                continue
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = RGBColor.from_string(color)
            ser.format.line.fill.background()
            if highlight and bars and i == first_main:
                base = th.neutral(300) if len(series) == 1 else mix(color, "FFFFFF", 0.6)
                for j, cat in enumerate(b["categories"]):
                    pt = ser.points[j]
                    pt.format.fill.solid()
                    pt.format.fill.fore_color.rgb = RGBColor.from_string(hl_color if str(cat) in highlight else base)
        if highlight and not (bars or pie):
            self.warnings.append(u"%s：highlight 只支援長條／橫條／圓餅圖，已忽略" % self._where)

        labels = b.get("labels") or pie
        if not pie:
            va, ca = chart.value_axis, chart.category_axis
            va.has_major_gridlines = not labels
            if va.has_major_gridlines:
                va.major_gridlines.format.line.color.rgb = RGBColor.from_string(th.neutral(100))
                va.major_gridlines.format.line.width = Pt(0.75)
            va.format.line.fill.background()
            va.major_tick_mark = XL_TICK_MARK.NONE
            va.tick_labels.font.color.rgb = RGBColor.from_string(th.neutral(500))
            va.tick_labels.font.size = Pt(th.size("chart"))
            if labels and bars:
                va.visible = False
            ca.format.line.color.rgb = RGBColor.from_string(th.neutral(300))
            ca.major_tick_mark = XL_TICK_MARK.NONE
            ca.tick_labels.font.size = Pt(th.size("chart"))
            if ctype in ("bar", "stacked_bar"):
                ca.reverse_order = True  # first category on top, as people read
            if pct:
                va.tick_labels.number_format = "0%"
                va.tick_labels.number_format_is_linked = False
            elif any(abs(v) >= 1000 for s in series for v in s["values"] if v is not None):
                va.tick_labels.number_format = "#,##0"
                va.tick_labels.number_format_is_linked = False
        if labels:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size = Pt(th.size("chart"))
            dl.font.color.rgb = RGBColor.from_string("FFFFFF" if pie or ctype in ("stacked", "stacked_bar") else th.neutral(700))
            big = any(abs(v) >= 1000 for s in series for v in s["values"] if v is not None)
            if pct or big:
                dl.number_format = "0%" if pct else "#,##0"
                dl.number_format_is_linked = False
            if ctype in ("column", "bar"):
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
        if bars:
            plot.gap_width = 60
            if len(series) > 1 and ctype in ("column", "bar"):
                plot.overlap = -10

    def _cell(self, cell, text, size, bold=False, emph=None, right=False, fill=None, color=None, rule=None) -> None:
        cell.margin_left = cell.margin_right = Emu(int(EMU_IN * 0.1))
        cell.margin_top = cell.margin_bottom = Emu(27432)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        if fill is not None:
            cell.fill.solid()
            _set_color(cell.fill.fore_color, fill)
        else:
            cell.fill.background()
        if rule:
            tcPr = cell._tc.get_or_add_tcPr()
            ln = etree.Element(qn("a:lnB"), w="9525")
            etree.SubElement(etree.SubElement(ln, qn("a:solidFill")), qn("a:srgbClr")).set("val", rule)
            tcPr.insert(0, ln)
        tf = cell.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        if right:
            p.alignment = PP_ALIGN.RIGHT
        self._runs(p, text, size, bold=bold, emph=emph, color=color)

    def _draw_table(self, slide, rect, b) -> None:
        cols, rows = b["columns"], b["rows"]
        if not cols:
            return
        if b.get("title"):
            rect = self._chart_title(slide, rect, b["title"])
        x, y, w, h = rect
        h_pt = h / float(EMU_PT)
        font = self.theme.size("table")

        def cap(f):
            return max(1, int(h_pt // (f * 2.2)) - 1)

        while font > 9 and len(rows) > cap(font):
            font = self.theme.snap(font - 0.5)
        if len(rows) > cap(font):
            self.warnings.append("%s：表格 %d 列放不下，只顯示前 %d 列（可用 limit 篩選或拆頁）"
                                 % (self._where, len(rows), cap(font)))
            rows = rows[:cap(font)]
        nrows = len(rows) + 1
        row_h = int(font * 2.2 * EMU_PT)
        gf = slide.shapes.add_table(nrows, len(cols), Emu(x), Emu(y), Emu(w), Emu(min(h, row_h * nrows)))
        tbl = gf.table
        style_id = gf._element.graphic.graphicData.tbl.tblPr.find(qn("a:tableStyleId"))
        if style_id is not None:
            style_id.text = _NO_STYLE_TABLE
        tbl.first_row, tbl.horz_banding = True, False
        weights = []
        for j, c in enumerate(cols):
            texts = [plain_text(str(c))] + [plain_text(str(r[j])) if j < len(r) and r[j] is not None else "" for r in rows]
            weights.append(min(24.0, max(3.0, max(em_width(t) for t in texts))))
        total, used = sum(weights), 0
        for j, wt in enumerate(weights):
            cw = int(w * wt / total) if j < len(weights) - 1 else w - used
            tbl.columns[j].width = Emu(cw)
            used += cw
        for i in range(nrows):
            tbl.rows[i].height = Emu(row_h)
        numeric = []
        for j in range(len(cols)):
            vals = [str(r[j]) for r in rows if j < len(r) and r[j] not in (None, "")]
            numeric.append(bool(vals) and sum(1 for v in vals if _NUMERIC_CELL.match(v)) >= len(vals) * 0.6)
        for j, c in enumerate(cols):
            self._cell(tbl.cell(0, j), str(c), font, bold=True, right=numeric[j],
                       fill=MSO_THEME_COLOR.ACCENT_1, color="FFFFFF")
        highlight = set(b.get("highlight") or ())
        hl_fill = mix(self.theme.alert, "FFFFFF", 0.88)
        for i, r in enumerate(rows, 1):
            hl = (i - 1) in highlight
            fill = hl_fill if hl else (self.theme.neutral(50) if i % 2 == 0 else "FFFFFF")
            for j in range(len(cols)):
                v = r[j] if j < len(r) and r[j] is not None else ""
                v = str(v)
                self._cell(tbl.cell(i, j), v, font, bold=hl, emph="alert" if hl else None,
                           right=numeric[j] and bool(_NUMERIC_CELL.match(v)) or numeric[j] and not v,
                           fill=fill, color=self.theme.neutral(900), rule=self.theme.neutral(100))

    def _draw_image(self, slide, rect, b) -> None:
        from PIL import Image

        x, y, w, h = rect
        with Image.open(str(b["path"])) as im:
            iw, ih = im.size
        s = min(w / float(iw), h / float(ih))
        pw, ph = int(iw * s), int(ih * s)
        slide.shapes.add_picture(str(b["path"]), Emu(x + (w - pw) // 2), Emu(y + (h - ph) // 2), Emu(pw), Emu(ph))

    def _notes(self, slide, notes) -> None:
        if not notes:
            return
        tf = slide.notes_slide.notes_text_frame
        if tf is not None:
            tf.text = notes
