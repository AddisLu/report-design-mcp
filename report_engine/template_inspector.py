"""Read a template's layouts/placeholders into a layout catalog the model can choose from."""
from __future__ import annotations

from pathlib import Path

from pptx.enum.shapes import PP_PLACEHOLDER as P

from .config import read_json, write_json
from .paths import design_dir
from .template import open_presentation

TITLE_TYPES = {P.TITLE, P.CENTER_TITLE, P.VERTICAL_TITLE}
SUBTITLE_TYPES = {P.SUBTITLE}
SKIP_TYPES = {P.DATE, P.FOOTER, P.SLIDE_NUMBER, P.HEADER}
TEXT_KINDS = {"BODY", "OBJECT", "VERTICAL_BODY", "VERTICAL_OBJECT"}


def _grid(ph, sw, sh):
    return [round(ph.left / sw * 12, 1), round(ph.top / sh * 8, 1),
            round(ph.width / sw * 12, 1), round(ph.height / sh * 8, 1)]


def _capacity(ph, pt=18.0):
    w = ph.width / 12700.0 - 14.4
    h = ph.height / 12700.0 - 7.2
    return max(0, int((w / pt) * (h / (pt * 1.3))))


def _guess_use(title, subtitle, content, sh):
    n = len(content)
    if title is not None and subtitle is not None and n == 0:
        return "封面／章節頁（標題＋副標）"
    if title is not None and n == 0:
        return "只有標題（自訂排版的底版）"
    if title is None and n == 0:
        return "空白"
    if n == 1 and content[0].height < sh * 0.25:
        return "章節頁（標題＋短說明）"
    if n == 1:
        return "標題＋單一內容區（條列、圖表或表格）"
    if n == 2:
        return "標題＋兩個內容區（比較、文字＋圖表）"
    return "標題＋%d 個內容區" % n


def _pick(layouts, pred, prefer):
    cands = [L for L in layouts if pred(L)]
    for word in prefer:
        for L in cands:
            if word in L["name"].lower():
                return L["key"]
    cands.sort(key=lambda L: -L["_area"])
    return cands[0]["key"] if cands else None


def inspect_presentation(prs) -> dict:
    sw, sh = prs.slide_width, prs.slide_height
    layouts, seen = [], {}
    for mi, master in enumerate(prs.slide_masters):
        for li, layout in enumerate(master.slide_layouts):
            name = (layout.name or "Layout %d" % (li + 1)).strip()
            seen[name] = seen.get(name, 0) + 1
            key = name if seen[name] == 1 else "%s #%d" % (name, seen[name])
            title = subtitle = None
            content = []
            for ph in layout.placeholders:
                try:
                    t = ph.placeholder_format.type
                    ok = ph.left is not None and ph.width is not None
                except Exception:
                    continue
                if t in SKIP_TYPES or not ok:
                    continue
                if t in TITLE_TYPES:
                    title = title or ph
                elif t in SUBTITLE_TYPES:
                    subtitle = subtitle or ph
                else:
                    content.append(ph)
            band = sh / 4.0
            content.sort(key=lambda p: (int(p.top // band), p.left))
            layouts.append({
                "key": key,
                "name": name,
                "master": mi,
                "index": li,
                "use": _guess_use(title, subtitle, content, sh),
                "title": title is not None,
                "subtitle": subtitle is not None,
                "content": [{
                    "kind": getattr(p.placeholder_format.type, "name", str(p.placeholder_format.type)),
                    "grid": _grid(p, sw, sh),
                    "capacity_chars": _capacity(p),
                } for p in content],
                "obj": layout,
                "title_idx": title.placeholder_format.idx if title is not None else None,
                "subtitle_idx": subtitle.placeholder_format.idx if subtitle is not None else None,
                "content_idx": [p.placeholder_format.idx for p in content],
                "title_bottom": int(title.top + title.height) if title is not None else None,
                "_area": sum(int(p.width) * int(p.height) for p in content),
            })

    def nc(L):
        return len(L["content_idx"])

    with_title = (
        _pick(layouts, lambda L: L["title"] and not L["subtitle"] and nc(L) == 0,
              ["title only", "僅標題", "只有標題"])
        or _pick(layouts, lambda L: L["title"], ["title and content", "標題及內容"])
        or layouts[0]["key"]
    )
    blank = _pick(layouts, lambda L: not L["title"] and not L["subtitle"] and nc(L) == 0, ["blank", "空白"]) or with_title
    cover = _pick(layouts, lambda L: L["title"] and L["subtitle"] and nc(L) == 0, ["title slide", "標題投影片", "封面"])
    section = None
    for L in layouts:  # only a layout that is clearly named as a section divider
        if L["key"] != cover and L["title"] and nc(L) <= 1 and any(
                w in L["name"].lower() for w in ("section", u"章節", u"節標題", u"區段")):
            section = L["key"]
            break
    by_count = {}
    for n in range(1, 7):
        k = _pick(layouts, lambda L, n=n: L["title"] and not L["subtitle"] and nc(L) == n,
                  ["title and content", "標題及內容", "two content", "兩個內容", "比較", "comparison"])
        if k:
            by_count[n] = k
    return {"layouts": layouts, "with_title": with_title, "blank": blank, "cover": cover, "section": section,
            "by_count": by_count,
            "slide_size": (int(sw), int(sh))}


def catalog_json(info: dict, template_path: Path, mtime: int) -> dict:
    sw, sh = info["slide_size"]
    return {
        "template": str(template_path),
        "template_mtime": mtime,
        "slide_size_in": [round(sw / 914400.0, 2), round(sh / 914400.0, 2)],
        "layouts": [{k: v for k, v in L.items() if k in ("key", "name", "use", "title", "subtitle", "content")}
                    for L in info["layouts"]],
        "custom_base": {"with_title": info["with_title"], "blank": info["blank"]},
        "cover": info["cover"],
        "section": info.get("section"),
        "by_count": {str(k): v for k, v in info["by_count"].items()},
    }


def ensure_catalog(template_path, force: bool = False) -> dict:
    """Regenerate design_system/layout_catalog.json when the template changes; keep hand-edited "use" notes."""
    path = design_dir() / "layout_catalog.json"
    tp = Path(template_path)
    mtime = int(tp.stat().st_mtime)
    old = read_json(path, {})
    if not force and old.get("template") == str(tp) and old.get("template_mtime") == mtime:
        return old
    cat = catalog_json(inspect_presentation(open_presentation(tp)), tp, mtime)
    if old.get("template") == str(tp):
        old_use = {L.get("key"): L.get("use") for L in old.get("layouts", [])}
        for L in cat["layouts"]:
            if old_use.get(L["key"]):
                L["use"] = old_use[L["key"]]
    write_json(path, cat)
    return cat
