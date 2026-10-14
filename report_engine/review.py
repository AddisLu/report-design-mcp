"""review_deck: the final art-direction pass over the whole deck — palette discipline, fonts, alignment,
rhythm, and a thumbnail overview of every slide (PowerPoint on Windows)."""
from __future__ import annotations

import base64
import io
from collections import Counter, defaultdict

from PIL import Image, ImageDraw
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn

from .renderer import EMU_IN
from .template import open_presentation
from .template_inspector import TITLE_TYPES
from .theme import load_theme, rgb_tuple

CHROME = (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE)
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _seg_dist(c, p, q):
    """Distance in RGB from colour c to the segment p→q (tints/shades of a palette colour count as on-palette)."""
    import math
    d = [b - a for a, b in zip(p, q)]
    L = sum(x * x for x in d) or 1.0
    t = max(0.0, min(1.0, sum((ci - pi) * di for ci, pi, di in zip(c, p, d)) / L))
    return math.sqrt(sum((ci - (pi + t * di)) ** 2 for ci, pi, di in zip(c, p, d)))


def _palette(theme):
    from .images import DOTS

    cols = {theme.primary, theme.secondary, theme.alert, "FFFFFF", "000000"} | set(theme.chart_colors) | set(DOTS)
    cols |= {theme.neutral(k) for k in ("900", "700", "500", "300", "100", "50")}
    for k in ("good", "warn", "bad"):
        cols |= {theme.status(k), theme.status_text(k)}
    return [rgb_tuple(c) for c in cols]


def _on_palette(hex6, pal):
    c = rgb_tuple(hex6)
    return any(min(_seg_dist(c, p, (255, 255, 255)), _seg_dist(c, p, (0, 0, 0))) < 30 for p in pal)


def _walk(shapes):
    for shp in shapes:
        yield shp
        if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
            for sub in _walk(shp.shapes):
                yield sub


def _kind(shp):
    if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
        return shp.name.split()[0]
    if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "Picture"
    if getattr(shp, "has_chart", False) and shp.has_chart:
        return "Chart"
    if getattr(shp, "has_table", False) and shp.has_table:
        return "Table"
    if shp.has_text_frame and shp.text_frame.text.strip():
        return "Text"
    return None


def _ph_type(shp):
    try:
        return shp.placeholder_format.type if shp.is_placeholder else None
    except Exception:
        return None


def overview_sheet(pptx_path, total, cols=4):
    from . import preview

    imgs, err = preview.export_images(pptx_path, "1-%d" % min(total, 30), total, max_edge=480, limit=30)
    if not imgs:
        return None, err
    tiles = []
    for it in imgs:
        tiles.append((it["slide"], Image.open(io.BytesIO(base64.b64decode(it["data"]))).convert("RGB")))
    tw, th = tiles[0][1].size
    rows = -(-len(tiles) // cols)
    sheet = Image.new("RGB", (cols * (tw + 8) + 8, rows * (th + 8) + 8), "#9AA3AD")
    d = ImageDraw.Draw(sheet)
    for i, (n, im) in enumerate(tiles):
        x, y = 8 + (i % cols) * (tw + 8), 8 + (i // cols) * (th + 8)
        sheet.paste(im, (x, y))
        d.rectangle([x, y, x + 22, y + 14], fill="#1F2933")
        d.text((x + 4, y + 2), str(n), fill="white")
    buf = io.BytesIO()
    sheet.save(buf, format="JPEG", quality=80)
    return base64.b64encode(buf.getvalue()).decode("ascii"), ""


def review_deck(path, cfg=None) -> dict:
    prs = open_presentation(path)
    theme = load_theme(cfg)
    pal = _palette(theme)
    fonts_ok = {theme.font_latin, theme.font_ea, "Arial"}
    issues = []
    off_colors = defaultdict(set)
    off_fonts = defaultdict(set)
    title_left = {}
    sigs = []
    text_only = []
    busy = []
    for n, slide in enumerate(prs.slides, 1):
        divider = slide.slide_layout._element.get("showMasterSp") == "0"
        accents = set()
        for shp in _walk(slide.shapes):
            if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
                continue
            el = shp._element
            for clr in el.iter(A + "srgbClr"):
                parent = clr.getparent()
                if parent is not None and parent.tag in (A + "outerShdw", A + "innerShdw", A + "glow"):
                    continue
                v = (clr.get("val") or "").upper()
                if len(v) != 6:
                    continue
                if not _on_palette(v, pal):
                    off_colors[v].add(n)
                c = rgb_tuple(v)
                if max(c) - min(c) > 40:  # chromatic, i.e. not a grey
                    accents.add(v)
            for tag in ("latin", "ea"):
                for f in el.iter(A + tag):
                    face = f.get("typeface") or ""
                    if face and not face.startswith("+") and face not in fonts_ok:
                        off_fonts[face].add(n)
        if len(accents) > 5 and not divider:
            busy.append((n, len(accents)))
        kinds = []
        body_left = []
        for shp in slide.shapes:
            t = _ph_type(shp)
            if t in CHROME:
                continue
            if t in TITLE_TYPES and shp.left is not None:
                if not divider:
                    title_left[n] = int(shp.left)
                continue
            k = _kind(shp)
            if k and shp.left is not None and not (shp.top > prs.slide_height - EMU_IN * 0.55 and shp.height < EMU_IN * 0.45):
                kinds.append(k)
                body_left.append(int(shp.left))
        if divider:
            sigs.append(None)
            continue
        sigs.append(tuple(sorted(Counter(kinds).items())))
        if kinds and all(k == "Text" for k in kinds):
            text_only.append(n)
        if n in title_left and body_left and abs(min(body_left) - title_left[n]) > EMU_IN * 0.15:
            issues.append({"slide": n, "level": "info", "msg": u"內容左緣和標題左緣差 %.2f 吋，沒有對齊同一條線"
                           % (abs(min(body_left) - title_left[n]) / float(EMU_IN))})
    if off_colors:
        for v, slides in sorted(off_colors.items(), key=lambda kv: -len(kv[1]))[:8]:
            issues.append({"slide": 0, "level": "warn", "msg": u"顏色 #%s 不在品牌色票內（投影片 %s）：改用 primary／neutral／狀態色"
                           % (v, u"、".join(str(s) for s in sorted(slides)))})
    for face, slides in off_fonts.items():
        issues.append({"slide": 0, "level": "warn", "msg": u"字型「%s」不是設計系統字型（投影片 %s）"
                       % (face, u"、".join(str(s) for s in sorted(slides)))})
    if title_left:
        mode = Counter(title_left.values()).most_common(1)[0][0]
        odd = [s for s, v in title_left.items() if abs(v - mode) > EMU_IN * 0.1]
        if odd:
            issues.append({"slide": 0, "level": "warn", "msg": u"標題位置不一致（投影片 %s）" % u"、".join(str(s) for s in odd)})
    run = 1
    for i in range(1, len(sigs) + 1):
        if i < len(sigs) and sigs[i] is not None and sigs[i] == sigs[i - 1]:
            run += 1
            continue
        if run >= 3:
            issues.append({"slide": 0, "level": "info", "msg": u"投影片 %d–%d 連續 %d 頁版面相同，節奏單調：換個元件或排法"
                           % (i - run + 1, i, run)})
        run = 1
    content = [s for s in sigs if s is not None]
    if content and len(text_only) > max(2, len(content) * 0.4):
        issues.append({"slide": 0, "level": "info", "msg": u"純文字頁偏多（投影片 %s）：把數字改成 KPI、圖表或時程"
                       % u"、".join(str(s) for s in text_only)})
    for n, k in busy:
        issues.append({"slide": n, "level": "warn", "msg": u"用了 %d 種彩色，太花：只留主色＋狀態色" % k})
    return {"slide_count": len(prs.slides), "issues": issues}


def format_review(res) -> str:
    issues = res["issues"]
    if not issues:
        return u"美術總檢：%d 頁，配色、字型、對齊、節奏都一致。" % res["slide_count"]
    warns = sum(1 for i in issues if i["level"] == "warn")
    lines = [u"美術總檢：%d 頁，%d 個需要修正、%d 個建議" % (res["slide_count"], warns, len(issues) - warns)]
    for i in sorted(issues, key=lambda x: (x["level"] != "warn", x["slide"])):
        where = u"整份" if i["slide"] == 0 else u"投影片 %d" % i["slide"]
        lines.append(u"- %s %s：%s" % (u"⚠" if i["level"] == "warn" else u"💡", where, i["msg"]))
    return "\n".join(lines)
