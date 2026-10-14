"""Review loop: build fingerprints, manual-edit detection and PowerPoint comments (old and new style)."""
from __future__ import annotations

import datetime as _dt
import difflib
import hashlib

from lxml import etree
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from .template import open_presentation

RT_MODERN_COMMENTS = "http://schemas.microsoft.com/office/2018/10/relationships/comments"
RT_AUTHORS = "http://schemas.microsoft.com/office/2018/10/relationships/authors"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_P188 = "http://schemas.microsoft.com/office/powerpoint/2018/8/main"
PREVIEW_CHARS = 1500


def _walk(shapes):
    for shp in shapes:
        yield shp
        if shp.shape_type == 6:  # group
            for sub in _walk(shp.shapes):
                yield sub


def slide_texts(prs):
    """Visible text of every slide (groups and tables included), in reading order of the shape tree."""
    out = []
    for slide in prs.slides:
        parts = []
        for shp in _walk(slide.shapes):
            if shp.has_text_frame and shp.text_frame.text.strip():
                parts.append(shp.text_frame.text.strip())
            if getattr(shp, "has_table", False) and shp.has_table:
                parts.extend(c.text.strip() for r in shp.table.rows for c in r.cells if c.text.strip())
        out.append("\n".join(parts))
    return out


def _sha1(data) -> str:
    return hashlib.sha1(data if isinstance(data, bytes) else data.encode("utf-8")).hexdigest()


def fingerprint(path, mapping):
    """What build_deck remembers about the file it wrote (stored as spec["_build"])."""
    with open(str(path), "rb") as f:
        whole = _sha1(f.read())
    texts = slide_texts(open_presentation(path))
    return {"output": str(path), "built_at": _dt.datetime.now().isoformat(timespec="seconds"), "sha1": whole,
            "mapping": mapping, "slides": [{"sha1": _sha1(t), "text": t[:PREVIEW_CHARS]} for t in texts]}


def slide_to_spec(build_info):
    out = {}
    for m in (build_info or {}).get("mapping") or []:
        for n in m.get("slides") or []:
            out[n] = m["spec"]
    return out


def manual_edits(build_info, path):
    """Slides whose text changed since build_deck wrote the file → [{slide, spec, before, after}]."""
    if not build_info or not path:
        return []
    try:
        with open(str(path), "rb") as f:
            if _sha1(f.read()) == build_info.get("sha1"):
                return []
    except OSError:
        return []
    before = build_info.get("slides") or []
    texts = slide_texts(open_presentation(path))
    spec_of = slide_to_spec(build_info)
    edits = []
    for i, t in enumerate(texts):
        old = before[i] if i < len(before) else None
        if old is None or _sha1(t) != old.get("sha1"):
            edits.append({"slide": i + 1, "spec": spec_of.get(i + 1), "before": (old or {}).get("text", u"（新增的投影片）"),
                          "after": t[:PREVIEW_CHARS]})
    if len(texts) < len(before):
        for i in range(len(texts), len(before)):
            edits.append({"slide": i + 1, "spec": spec_of.get(i + 1), "before": before[i].get("text", ""),
                          "after": u"（這張投影片被刪除）"})
    if not edits:  # same text, other changes (moved shapes, colours…)
        edits.append({"slide": 0, "spec": None, "before": "", "after": u"（文字沒變，但檔案在 PowerPoint 裡改過：可能調整了位置、顏色或圖片）"})
    return edits


def format_edits(edits) -> str:
    if not edits:
        return ""
    lines = [u"### ✋ PowerPoint 裡手動改過的內容（重新產生會蓋掉，請把要保留的改動寫回 spec）"]
    for e in edits:
        if not e["slide"]:
            lines.append(u"- " + e["after"])
            continue
        where = u"投影片 %d%s" % (e["slide"], u"（spec 第 %d 頁）" % e["spec"] if e.get("spec") else "")
        old, new = e["before"].split("\n"), e["after"].split("\n")
        removed, added = [], []
        for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, old, new).get_opcodes():
            if op in ("replace", "delete"):
                removed += old[i1:i2]
            if op in ("replace", "insert"):
                added += new[j1:j2]
        lines.append(u"- " + where)
        for t in removed[:6]:
            lines.append(u"  - 原本：" + t[:120])
        for t in added[:6]:
            lines.append(u"  - 改成：" + t[:120])
    return "\n".join(lines)


def _xml(part):
    return etree.fromstring(part.blob)


def _authors(prs):
    names = {}
    for rel in prs.part.rels.values():
        if rel.is_external:
            continue
        if rel.reltype == RT.COMMENT_AUTHORS:
            for a in _xml(rel.target_part).iter("{%s}cmAuthor" % NS_P):
                names[("old", a.get("id"))] = a.get("name") or a.get("initials") or "?"
        elif rel.reltype == RT_AUTHORS:
            for a in _xml(rel.target_part).iter("{%s}author" % NS_P188):
                names[("new", a.get("id"))] = a.get("name") or a.get("initials") or "?"
    return names


def _text(el):
    return "".join(t.text or "" for t in el.iter("{%s}t" % NS_A)).strip()


def read_comments(path):
    """-> [{slide, author, text, resolved, replies: [(author, text)]}] for old- and new-style comments."""
    prs = open_presentation(path)
    authors = _authors(prs)
    out = []
    for n, slide in enumerate(prs.slides, 1):
        for rel in slide.part.rels.values():
            if rel.is_external:
                continue
            if rel.reltype == RT.COMMENTS:
                for cm in _xml(rel.target_part).iter("{%s}cm" % NS_P):
                    t = cm.find("{%s}text" % NS_P)
                    out.append({"slide": n, "author": authors.get(("old", cm.get("authorId")), "?"),
                                "text": (t.text or "").strip() if t is not None else "", "resolved": False, "replies": []})
            elif rel.reltype == RT_MODERN_COMMENTS:
                root = _xml(rel.target_part)
                for cm in root.findall("{%s}cm" % NS_P188):
                    body = cm.find("{%s}txBody" % NS_P188)
                    replies = []
                    rl = cm.find("{%s}replyLst" % NS_P188)
                    if rl is not None:
                        for rp in rl.findall("{%s}reply" % NS_P188):
                            rb = rp.find("{%s}txBody" % NS_P188)
                            replies.append((authors.get(("new", rp.get("authorId")), "?"), _text(rb) if rb is not None else ""))
                    out.append({"slide": n, "author": authors.get(("new", cm.get("authorId")), "?"),
                                "text": _text(body) if body is not None else "",
                                "resolved": cm.get("status") in ("resolved", "closed"), "replies": replies})
    return out


def format_feedback(path, comments, build_info, edits, titles) -> str:
    spec_of = slide_to_spec(build_info)
    open_items = [c for c in comments if not c["resolved"]]
    lines = [u"# 審閱回饋：%s" % path,
             u"共 %d 則註解（未解決 %d 則）%s" % (len(comments), len(open_items),
                                         u"" if build_info else u"；找不到 .spec.json，只能列投影片編號")]
    if comments:
        lines += ["", u"| 投影片 | spec 頁 | 目前標題 | 誰 | 註解 | 狀態 |", "|---|---|---|---|---|---|"]
        for c in comments:
            reply = u"（回覆：%s）" % u"；".join(u"%s：%s" % r for r in c["replies"]) if c["replies"] else ""
            lines.append(u"| %d | %s | %s | %s | %s%s | %s |" % (
                c["slide"], spec_of.get(c["slide"], u"—"), (titles[c["slide"] - 1] if c["slide"] <= len(titles) else "")[:20],
                c["author"], c["text"].replace("|", u"／").replace("\n", " "), reply, u"已解決" if c["resolved"] else u"未解決"))
    else:
        lines.append(u"（沒有註解。請主管在 PowerPoint 用「校閱 → 新增註解」留言後再試。）")
    if edits:
        lines += ["", format_edits(edits)]
    lines += ["", u"**下一步**：把每則未解決的註解轉成修改，合併成一次 patch_deck（spec 頁碼見上表；會另存新版本、不覆蓋原檔），"
                  u"再用表格回覆使用者「哪則註解 → 怎麼改」。看不懂或需要決策的註解先問使用者。"]
    return "\n".join(lines)
