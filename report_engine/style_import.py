"""import_style: learn the brand palette, fonts and logo from an existing company deck → theme.yaml."""
from __future__ import annotations

import base64
import io
import shutil
from collections import Counter, defaultdict
from pathlib import Path

import yaml
from lxml import etree
from PIL import Image, ImageDraw
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from .config import _read_yaml
from .paths import design_dir
from .template import open_presentation
from .theme import DEFAULT_THEME, contrast, mix, rgb_tuple

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
SCHEME_TAGS = ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6")
SCHEME_ALIAS = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}


def _chromatic(hex6):
    r, g, b = rgb_tuple(hex6)
    return max(r, g, b) - min(r, g, b) > 40 and 40 < (r + g + b) / 3.0 < 235


def _hue_far(a, b):
    return sum(abs(x - y) for x, y in zip(rgb_tuple(a), rgb_tuple(b))) > 120


def _theme_root(prs):
    part = prs.slide_masters[0].part.part_related_by(RT.THEME)
    el = getattr(part, "_element", None)
    return el if el is not None else etree.fromstring(part.blob)


def _scheme(root):
    out = {}
    cs = root.find(".//%sclrScheme" % A)
    if cs is None:
        return out
    for tag in SCHEME_TAGS:
        el = cs.find(A + tag)
        if el is None or not len(el):
            continue
        c = el[0]
        val = c.get("val") if c.tag == A + "srgbClr" else c.get("lastClr")
        if val and len(val) == 6:
            out[tag] = val.upper()
    return out


def _fonts(root):
    fs = root.find(".//%sfontScheme" % A)
    out = {}
    if fs is None:
        return out
    for which in ("minorFont", "majorFont"):
        f = fs.find(A + which)
        if f is None:
            continue
        latin = f.find(A + "latin")
        ea = f.find(A + "ea")
        hant = next((x.get("typeface") for x in f.findall(A + "font") if x.get("script") == "Hant"), None)
        if latin is not None and latin.get("typeface") and "latin" not in out:
            out["latin"] = latin.get("typeface")
        face = (ea.get("typeface") if ea is not None else "") or hant
        if face and "ea" not in out:
            out["ea"] = face
    return out


def _walk(shapes):
    for shp in shapes:
        yield shp
        if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
            for sub in _walk(shp.shapes):
                yield sub


P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def _fill_color(container, scheme):
    sf = container.find(A + "solidFill") if container is not None else None
    if sf is None or not len(sf):
        return None
    c = sf[0]
    if c.tag == A + "srgbClr":
        return (c.get("val") or "").upper()
    if c.tag == A + "schemeClr":
        return scheme.get(SCHEME_ALIAS.get(c.get("val"), c.get("val")))
    return None


def _usage(prs, scheme):
    """Area-weighted use of fill colours (shapes, table cells, master/layout backgrounds). Text colours are
    ignored: status words and alerts would otherwise outvote the brand colour."""
    use = Counter()
    page = int(prs.slide_width) * int(prs.slide_height) // 10 ** 8

    def shapes_fill(shapes, weight=1.0):
        for shp in _walk(shapes):
            if shp.shape_type == MSO_SHAPE_TYPE.PICTURE or shp.width is None:
                continue
            area = int(shp.width) * int(shp.height) // 10 ** 8
            c = _fill_color(shp._element.find(P + "spPr"), scheme)
            if c:
                use[c] += area * weight
            if getattr(shp, "has_table", False) and shp.has_table:
                cell_area = area // max(1, len(shp.table.rows) * len(shp.table.columns))
                for row in shp.table.rows:
                    for cell in row.cells:
                        c = _fill_color(cell._tc.tcPr, scheme)
                        if c:
                            use[c] += cell_area * weight

    master = prs.slide_masters[0]
    holders = [master] + list(master.slide_layouts)
    for h in holders:
        bg = h._element.find("%scSld/%sbg/%sbgPr" % (P, P, P))
        c = _fill_color(bg, scheme)
        if c:
            use[c] += page * 0.5
        shapes_fill(h.shapes, 2.0)
    for slide in prs.slides:
        shapes_fill(slide.shapes)
    a1 = scheme.get("accent1")
    if a1 and _chromatic(a1):  # company templates put the brand colour in accent1
        use[a1] += (max(use.values()) if use else 1) * 0.5
    return use


def _logo(prs):
    """(blob, ext, where) of the logo: a picture on the master/layouts, else one repeated on >=50% of slides."""
    for holder in [prs.slide_masters[0]] + list(prs.slide_masters[0].slide_layouts):
        for shp in holder.shapes:
            if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
                return shp.image.blob, shp.image.ext, u"母片／版型"
    seen = defaultdict(set)
    blobs = {}
    for n, slide in enumerate(prs.slides):
        for shp in slide.shapes:
            if shp.shape_type == MSO_SHAPE_TYPE.PICTURE and shp.width < prs.slide_width * 0.3:
                key = (shp.image.sha1, int(shp.left) // 200000, int(shp.top) // 200000)
                seen[key].add(n)
                blobs[key] = (shp.image.blob, shp.image.ext)
    total = len(prs.slides)
    if total >= 2 and seen:
        key, slides = max(seen.items(), key=lambda kv: len(kv[1]))
        if len(slides) >= max(2, total * 0.5):
            blob, ext = blobs[key]
            return blob, ext, u"%d／%d 張投影片的同一位置" % (len(slides), total)
    return None, None, ""


def _neutrals(dark):
    base = dark if dark and not _chromatic(dark) and contrast(dark, "FFFFFF") > 10 else DEFAULT_THEME["colors"]["neutral"]["900"]
    if base == "000000":
        base = "1F2933"
    n500 = mix(base, "FFFFFF", 0.42)
    t = 0.42
    while contrast(n500, "FFFFFF") < 4.6 and t > 0.1:
        t -= 0.04
        n500 = mix(base, "FFFFFF", t)
    return {"900": base, "700": mix(base, "FFFFFF", 0.2), "500": n500, "300": mix(base, "FFFFFF", 0.75),
            "100": mix(base, "FFFFFF", 0.88), "50": mix(base, "FFFFFF", 0.95)}


def extract_style(path):
    prs = open_presentation(path)
    root = _theme_root(prs)
    scheme = _scheme(root)
    fonts = _fonts(root)
    use = _usage(prs, scheme)
    status = set(DEFAULT_THEME["colors"]["status"].values())  # our own status colours are not brand colours
    ranked = [c for c, _n in use.most_common() if len(c) == 6 and _chromatic(c) and c not in status]
    accents = [scheme[k] for k in SCHEME_TAGS[4:] if k in scheme and _chromatic(scheme[k])]
    primary = ranked[0] if ranked else (accents[0] if accents else DEFAULT_THEME["colors"]["primary"])
    rest = [c for c in ranked[1:] + accents if c != primary]
    secondary = next((c for c in rest if _hue_far(c, primary)), rest[0] if rest else DEFAULT_THEME["colors"]["secondary"])
    chart = [primary, secondary]
    for c in rest + accents:
        if len(chart) >= 6:
            break
        if all(_hue_far(c, x) for x in chart):
            chart.append(c)
    chart += [c for c in DEFAULT_THEME["colors"]["chart"] if c not in chart][:max(0, 6 - len(chart))]
    notes = []
    if not ranked:
        notes.append(u"投影片裡幾乎沒有彩色填色，主色取自主題的 accent1")
    style = {"colors": {"primary": primary, "secondary": secondary, "chart": chart,
                        "neutral": _neutrals(scheme.get("dk1"))},
             "fonts": {"latin": fonts.get("latin") or DEFAULT_THEME["fonts"]["latin"],
                       "ea": fonts.get("ea") or DEFAULT_THEME["fonts"]["ea"]}}
    if not fonts.get("ea"):
        notes.append(u"範本沒有設定中文字型，沿用微軟正黑體")
    blob, ext, where = _logo(prs)
    return style, notes, (blob, ext, where)


def swatch(style, logo_blob=None):
    from .sketch import _font

    f, _ = _font(16)
    small, _ = _font(13)
    img = Image.new("RGB", (900, 300), "white")
    d = ImageDraw.Draw(img)
    c = style["colors"]
    row = [("primary", c["primary"]), ("secondary", c["secondary"])] + [("chart-%d" % (i + 3), x) for i, x in enumerate(c["chart"][2:6])]
    for i, (name, hx) in enumerate(row):
        x = 20 + i * 140
        d.rectangle([x, 20, x + 120, 110], fill="#" + hx)
        d.text((x, 116), name, fill="#333", font=small)
        d.text((x, 134), "#" + hx, fill="#666", font=small)
    for i, (k, hx) in enumerate(sorted(c["neutral"].items(), key=lambda kv: -int(kv[0]))):
        x = 20 + i * 100
        d.rectangle([x, 170, x + 90, 210], fill="#" + hx, outline="#ddd")
        d.text((x, 214), "neutral-%s" % k, fill="#666", font=small)
    d.text((20, 250), u"字型：%s ／ %s" % (style["fonts"]["latin"], style["fonts"]["ea"]), fill="#222", font=f)
    if logo_blob:
        try:
            lg = Image.open(io.BytesIO(logo_blob)).convert("RGBA")
            lg.thumbnail((200, 80))
            img.paste(lg, (680, 210), lg)
        except Exception:
            pass
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85)
    return base64.b64encode(buf.getvalue()).decode("ascii")


def apply_style(style, logo=None):
    """Merge into design_system/theme.yaml (backup kept as theme.yaml.bak) -> (theme path, logo path or None)."""
    path = design_dir() / "theme.yaml"
    current = _read_yaml(path) if path.exists() else {}
    if path.exists():
        shutil.copyfile(str(path), str(path.with_suffix(".yaml.bak")))
    merged = dict(current)
    colors = dict(current.get("colors") or {})
    colors.update(style["colors"])
    merged["colors"] = colors
    merged["fonts"] = dict(current.get("fonts") or {}, **style["fonts"])
    logo_path = None
    blob, ext = logo or (None, None)
    if blob:
        brand = design_dir() / "brand"
        brand.mkdir(parents=True, exist_ok=True)
        logo_file = brand / ("logo.%s" % (ext or "png"))
        logo_file.write_bytes(blob)
        try:
            rel = logo_file.relative_to(design_dir().parent)
            logo_path = str(rel).replace("\\", "/")
        except ValueError:
            logo_path = str(logo_file)
        merged["logo"] = dict(current.get("logo") or {}, path=logo_path)
    header = u"# 由 import_style 從公司簡報自動產生（原檔備份為 theme.yaml.bak），可再手動微調。\n"
    path.write_text(header + yaml.safe_dump(merged, allow_unicode=True, sort_keys=False), encoding="utf-8")
    return path, logo_path


def import_style(pptx_path, apply=False):
    """-> (text, swatch image b64)."""
    style, notes, (blob, ext, where) = extract_style(pptx_path)
    text = [u"# 從 %s 抽出的品牌設定" % Path(str(pptx_path)).name,
            "```yaml", yaml.safe_dump(style, allow_unicode=True, sort_keys=False).strip(), "```",
            u"Logo：%s" % (u"找到（%s）" % where if blob else u"沒找到，可手動在 theme.yaml 的 logo.path 指定")]
    text += [u"- " + n for n in notes]
    if apply:
        theme_path, logo_path = apply_style(style, (blob, ext) if blob else None)
        text.append(u"\n✅ 已寫入 %s（舊檔備份為 theme.yaml.bak）%s；下次 build_deck 會自動重建品牌範本。"
                    % (theme_path, u"，Logo 存到 %s" % logo_path if logo_path else ""))
    else:
        text.append(u"\n這是預覽，還沒寫入。使用者確認色票後，用 apply=true 再呼叫一次。")
    text.append(u"若這份簡報本身的版型就很好，也可以直接當公司範本：inspect_template（set_default=true）。")
    return "\n".join(text), swatch(style, blob)
