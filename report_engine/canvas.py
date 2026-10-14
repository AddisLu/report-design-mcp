"""Native PowerPoint drawing primitives shared by the SVG converter and the diagram blocks (gantt / timeline / process).

Everything is drawn as editable shapes inside one group, coloured from the theme, with text snapped to the type scale.
Coordinates are absolute EMU.
"""
from __future__ import annotations

import re

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

EMU_PT = 12700
_HEX3 = re.compile(r"^#?([0-9a-fA-F]{3})$")
_HEX6 = re.compile(r"^#?([0-9a-fA-F]{6})$")
_RGB = re.compile(r"^rgba?\(\s*([\d.]+%?)\s*,\s*([\d.]+%?)\s*,\s*([\d.]+%?)")
_VAR = re.compile(r"^var\(\s*--([\w-]+)\s*(?:,[^)]*)?\)$")
NAMED = {
    "black": "000000", "white": "FFFFFF", "red": "FF0000", "green": "008000", "blue": "0000FF",
    "gray": "808080", "grey": "808080", "orange": "FFA500", "yellow": "FFFF00", "purple": "800080",
    "navy": "000080", "teal": "008080", "silver": "C0C0C0", "lightgray": "D3D3D3", "lightgrey": "D3D3D3",
    "darkgray": "A9A9A9", "darkgrey": "A9A9A9",
}


def resolve_color(value, theme):
    """Theme token (primary, neutral-300, bad, chart-2…), hex, rgb(), or a basic name → 6-digit hex; None for none."""
    if value is None:
        return None
    s = str(value).strip()
    if not s or s.lower() in ("none", "transparent"):
        return None
    m = _VAR.match(s)
    if m:
        s = m.group(1)
    key = s.lower().replace("_", "-")
    tokens = {"primary": theme.primary, "accent": theme.primary, "secondary": theme.secondary,
              "alert": theme.alert, "text": theme.neutral(900), "muted": theme.neutral(500),
              "surface": theme.neutral(50), "line": theme.neutral(300)}
    if key in tokens:
        return tokens[key]
    if key in ("good", "warn", "bad"):
        return theme.status(key)
    if key.startswith(("status-", "text-")) and key.split("-", 1)[1] in ("good", "warn", "bad"):
        k = key.split("-", 1)[1]
        return theme.status_text(k) if key.startswith("text-") else theme.status(k)
    if key.startswith(("neutral-", "gray-", "grey-")):
        return theme.neutral(key.split("-", 1)[1])
    if key.startswith("chart-"):
        try:
            cs = theme.chart_colors
            return cs[(int(key.split("-")[1]) - 1) % len(cs)]
        except ValueError:
            return None
    m = _HEX6.match(s)
    if m:
        return m.group(1).upper()
    m = _HEX3.match(s)
    if m:
        return "".join(c * 2 for c in m.group(1)).upper()
    m = _RGB.match(s)
    if m:
        vals = []
        for v in m.groups():
            n = float(v[:-1]) * 2.55 if v.endswith("%") else float(v)
            vals.append(max(0, min(255, int(round(n)))))
        return "%02X%02X%02X" % tuple(vals)
    return NAMED.get(key)


def set_alt(shape, text):
    """Alternative text (cNvPr@descr) for screen readers, search and PowerPoint's accessibility checker."""
    from .renderer import plain_text

    t = " ".join(plain_text(str(text or "")).split())
    if shape is not None and t:
        shape._element._nvXxPr.cNvPr.set("descr", t[:600])


def _alpha(color_el, alpha):
    if alpha is None or alpha >= 0.999 or color_el is None:
        return
    a = etree.SubElement(color_el, qn("a:alpha"))
    a.set("val", str(int(max(0.0, alpha) * 100000)))


class Canvas:
    def __init__(self, renderer, slide, name="Diagram"):
        self.r = renderer
        self.theme = renderer.theme
        self.group = slide.shapes.add_group_shape()
        self.group.name = name
        self.shapes = self.group.shapes
        self.count = 0

    # -- styling ----------------------------------------------------------
    def _style(self, shp, fill=None, stroke=None, width=1.0, dash=False, alpha=None, stroke_alpha=None):
        style = shp._element.find(qn("p:style"))
        if style is not None:
            shp._element.remove(style)
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = RGBColor.from_string(fill)
            _alpha(shp._element.spPr.find("%s/%s" % (qn("a:solidFill"), qn("a:srgbClr"))), alpha)
        else:
            shp.fill.background()
        self._line(shp.line, stroke, width, dash, stroke_alpha)
        self.count += 1
        return shp

    @staticmethod
    def _line(line, stroke, width, dash, alpha=None):
        if stroke:
            line.color.rgb = RGBColor.from_string(stroke)
            line.width = Pt(max(0.25, width))
            if dash:
                line.dash_style = MSO_LINE.DASH
            ln = line._get_or_add_ln()
            _alpha(ln.find("%s/%s" % (qn("a:solidFill"), qn("a:srgbClr"))), alpha)
        else:
            line.fill.background()

    @staticmethod
    def _arrows(line, start=False, end=False):
        if not (start or end):
            return
        ln = line._get_or_add_ln()
        for tag, on in (("a:headEnd", start), ("a:tailEnd", end)):
            if on:
                el = etree.SubElement(ln, qn(tag))
                el.set("type", "triangle")
                el.set("w", "med")
                el.set("len", "med")

    # -- primitives -------------------------------------------------------
    def rect(self, x, y, w, h, fill=None, stroke=None, width=1.0, radius=0, dash=False, alpha=None, shape=None):
        if w <= 0 or h <= 0:
            return None
        kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius > 0 else MSO_SHAPE.RECTANGLE)
        shp = self.shapes.add_shape(kind, Emu(int(x)), Emu(int(y)), Emu(max(1, int(w))), Emu(max(1, int(h))))
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            shp.adjustments[0] = min(0.5, float(radius) / max(1.0, min(w, h)))
        return self._style(shp, fill, stroke, width, dash, alpha)

    def ellipse(self, cx, cy, rx, ry, fill=None, stroke=None, width=1.0, dash=False, alpha=None):
        if rx <= 0 or ry <= 0:
            return None
        shp = self.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(cx - rx)), Emu(int(cy - ry)), Emu(int(2 * rx)), Emu(int(2 * ry)))
        return self._style(shp, fill, stroke, width, dash, alpha)

    def line(self, x1, y1, x2, y2, stroke, width=1.0, dash=False, arrow_start=False, arrow_end=False, alpha=None):
        if not stroke:
            return None
        c = self.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(int(x1)), Emu(int(y1)), Emu(int(x2)), Emu(int(y2)))
        style = c._element.find(qn("p:style"))
        if style is not None:
            c._element.remove(style)
        self._line(c.line, stroke, width, dash, alpha)
        self._arrows(c.line, arrow_start, arrow_end)
        self.count += 1
        return c

    def poly(self, subpaths, fill=None, stroke=None, width=1.0, dash=False, alpha=None, stroke_alpha=None,
             arrow_start=False, arrow_end=False):
        """subpaths: [(points [(x, y)…], closed)] in EMU."""
        subpaths = [(pts, closed) for pts, closed in subpaths if len(pts) >= 2]
        if not subpaths or not (fill or stroke):
            return None
        first = subpaths[0][0][0]
        fb = self.shapes.build_freeform(first[0], first[1], scale=1.0)
        for i, (pts, closed) in enumerate(subpaths):
            if i:
                fb.move_to(pts[0][0], pts[0][1])
            fb.add_line_segments(pts[1:], close=closed and bool(fill))
        shp = fb.convert_to_shape()
        self._style(shp, fill, stroke, width, dash, alpha, stroke_alpha)
        self._arrows(shp.line, arrow_start, arrow_end)
        return shp

    def autoshape(self, kind, x, y, w, h, fill=None, stroke=None, width=1.0, adjust=None):
        shp = self.shapes.add_shape(kind, Emu(int(x)), Emu(int(y)), Emu(max(1, int(w))), Emu(max(1, int(h))))
        if adjust is not None:
            shp.adjustments[0] = adjust
        return self._style(shp, fill, stroke, width)

    def text(self, x, y, text, size, color=None, bold=False, align="start", valign="top", width=None,
             height=None, wrap=False, shape=None):
        """Place text. x is the start / centre / end per align; y is top / middle / baseline per valign.

        With `shape`, the text goes inside that shape (centred) instead of a new textbox.
        """
        from .renderer import em_width, plain_text

        size = self.theme.snap(size)
        if shape is not None:
            tf = shape.text_frame
            tf.margin_left = tf.margin_right = Emu(int(EMU_PT * 4))
            tf.margin_top = tf.margin_bottom = 0
            tf.word_wrap = wrap
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = {"start": PP_ALIGN.LEFT, "end": PP_ALIGN.RIGHT}.get(align, PP_ALIGN.CENTER)
            self.r._runs(p, text, size, bold=bold, color=color or self.theme.neutral(900))
            return shape
        lines = str(text).split("\n")
        est = max(em_width(plain_text(t)) for t in lines) * size * EMU_PT * 1.08 + EMU_PT * 2
        w = int(width or est)
        h = int(height or size * EMU_PT * (1.25 * len(lines)) + EMU_PT * 2)
        if align == "middle":
            left = x - w / 2.0
        elif align == "end":
            left = x - w
        else:
            left = x
        if valign == "middle":
            top = y - h / 2.0
        elif valign == "baseline":
            top = y - size * EMU_PT * 0.95
        elif valign == "bottom":
            top = y - h
        else:
            top = y
        tb = self.shapes.add_textbox(Emu(int(left)), Emu(int(top)), Emu(max(1, w)), Emu(max(1, h)))
        tf = tb.text_frame
        tf.word_wrap = wrap
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}.get(valign, MSO_ANCHOR.TOP)
        for i, t in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"middle": PP_ALIGN.CENTER, "end": PP_ALIGN.RIGHT}.get(align, PP_ALIGN.LEFT)
            self.r._runs(p, t, size, bold=bold, color=color or self.theme.neutral(900))
        self.count += 1
        return tb
