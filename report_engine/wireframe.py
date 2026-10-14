"""Wireframes of the planned slides (drawn from a trial build), side by side with the hand-drawn sketch page,
so the user can confirm how the sketch and its red-pen notes were read before any real work starts."""
from __future__ import annotations

import base64
import io
import re

from PIL import Image, ImageDraw
from pptx.enum.shapes import PP_PLACEHOLDER

from .review import _kind, _ph_type
from .sketch import _font, load_sketch_pages
from .template import open_presentation
from .template_inspector import TITLE_TYPES

W, H = 960, 540
CHROME = (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE)
LABELS = {
    "Text": (u"文字", (37, 99, 235)), "Chart": (u"圖表", (30, 78, 140)), "Table": (u"表格", (30, 78, 140)),
    "Picture": (u"圖片", (22, 128, 61)), "KPI": ("KPI", (30, 78, 140)), "Status": (u"狀態燈", (30, 78, 140)),
    "Gantt": (u"甘特圖", (124, 58, 237)), "Timeline": (u"時程表", (124, 58, 237)), "Process": (u"流程", (124, 58, 237)),
    "SVG": (u"示意圖", (124, 58, 237)), "Image": (u"截圖／圖片", (22, 128, 61)),
    "Placeholder": (u"⚠ 待補圖片", (217, 119, 6)),
}


def _snippet(shp, n=16):
    try:
        t = shp.text_frame.text.strip().split("\n")[0] if shp.has_text_frame else ""
    except Exception:
        t = ""
    return t[:n] + (u"…" if len(t) > n else "")


def _wire(slide, sw, sh, header, notes):
    img = Image.new("RGB", (W, H + 30 + 24 * max(1, min(len(notes), 4))), "white")
    d = ImageDraw.Draw(img)
    f_head, _ = _font(18)
    f_lab, _ = _font(15)
    f_small, _ = _font(13)
    d.rectangle([0, 0, W, 30], fill=(31, 41, 51))
    d.text((10, 5), header, fill="white", font=f_head)
    top = 30
    for i in range(1, 12):
        x = int(W * i / 12.0)
        d.line([(x, top), (x, top + H)], fill=(236, 239, 243))
    for j in range(1, 8):
        y = top + int(H * j / 8.0)
        d.line([(0, y), (W, y)], fill=(236, 239, 243))
    d.rectangle([0, top, W - 1, top + H], outline=(160, 168, 178))
    k = W / float(sw)
    for shp in slide.shapes:
        t = _ph_type(shp)
        if t in CHROME or shp.left is None or shp.width is None:
            continue
        x0, y0 = int(shp.left * k), top + int(shp.top * k)
        x1, y1 = int((shp.left + shp.width) * k), top + int((shp.top + shp.height) * k)
        if t in TITLE_TYPES:
            label, color = u"標題：" + _snippet(shp, 24), (31, 41, 51)
        else:
            kind = _kind(shp)
            if kind is None:
                continue
            name, color = LABELS.get(kind, (kind, (97, 110, 124)))
            snip = _snippet(shp) if kind == "Text" else ""
            label = name + (u"：" + snip if snip else "")
        d.rectangle([x0, y0, x1, y1], outline=color, width=3)
        tw = int(d.textlength(label, font=f_lab)) if hasattr(d, "textlength") else len(label) * 15
        d.rectangle([x0 + 3, y0 + 3, x0 + 9 + tw, y0 + 23], fill=color)
        d.text((x0 + 6, y0 + 4), label, fill="white", font=f_lab)
    y = top + H + 6
    if not notes:
        d.text((10, y), u"（這頁沒有手寫註解）", fill=(120, 120, 120), font=f_small)
    for raw, applied in notes[:4]:
        d.text((10, y), (u"註：%s → %s" % (raw, applied))[:70], fill=(198, 40, 40), font=f_small)
        y += 24
    return img


def _page_no(v):
    m = re.search(r"\d+", str(v or ""))
    return int(m.group(0)) if m else None


def wireframes(pptx_path, spec, mapping, slides=None, limit=8):
    """-> list of base64 JPEGs: [sketch page | wireframe] per slide (or the wireframe alone)."""
    prs = open_presentation(pptx_path)
    sw, sh = int(prs.slide_width), int(prs.slide_height)
    spec_of = {}
    for m in mapping:
        for n in m["slides"]:
            spec_of[n] = m["spec"]
    sketch = spec.get("sketch")
    sketch_pages = {}
    if sketch:
        try:
            for pg in load_sketch_pages(sketch):
                sketch_pages[pg["page"]] = pg["data"]
        except Exception:
            sketch_pages = {}
    wanted = set(slides) if slides else None
    out = []
    for n, slide in enumerate(prs.slides, 1):
        if wanted is not None and n not in wanted:
            continue
        if len(out) >= limit:
            break
        sidx = spec_of.get(n)
        s = spec["slides"][sidx - 1] if sidx and sidx <= len(spec.get("slides") or []) else {}
        notes = [(a.get("raw", ""), a.get("applied", "")) for a in (s.get("annotations") or []) if isinstance(a, dict)]
        page = _page_no(s.get("source_page"))
        header = u"投影片 %d（spec 第 %s 頁%s）" % (n, sidx or u"—", u"，草稿第 %d 頁" % page if page else "")
        wire = _wire(slide, sw, sh, header, notes)
        if page and page in sketch_pages:
            sk = Image.open(io.BytesIO(base64.b64decode(sketch_pages[page]))).convert("RGB")
            sk.thumbnail((W, wire.height))
            both = Image.new("RGB", (sk.width + 16 + wire.width, max(sk.height, wire.height)), (154, 163, 173))
            both.paste(sk, (0, 0))
            both.paste(wire, (sk.width + 16, 0))
            wire = both
        buf = io.BytesIO()
        wire.save(buf, format="JPEG", quality=82)
        out.append(base64.b64encode(buf.getvalue()).decode("ascii"))
    return out
