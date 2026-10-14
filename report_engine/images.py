"""Screenshot / photo finishing with native PowerPoint features (nothing is flattened, everything stays editable).

- style: card (rounded corners + hairline border + soft shadow), browser (window frame with URL bar),
  shadow, plain
- crop (non-destructive, fractions or pixels) and trim (auto-remove uniform borders around a screenshot)
- highlights: numbered boxes on the screenshot + a numbered legend; zoom: magnified inset of a region
- caption under the image; gallery grid for many images
"""
from __future__ import annotations

from lxml import etree
from PIL import Image, ImageChops, ImageDraw
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu

from .canvas import EMU_PT, Canvas, resolve_color

EMU_IN = 914400
DOTS = ("FF5F57", "FEBC2E", "28C840")


def trim_box(path, tolerance=12):
    """Pixel bbox (l, t, r, b) of the content inside a uniform border, or None."""
    with Image.open(str(path)) as im:
        rgb = im.convert("RGB")
        w, h = rgb.size
        bg = rgb.getpixel((0, 0))
        if any(abs(a - b) > tolerance for a, b in zip(rgb.getpixel((w - 1, h - 1)), bg)):
            return None
        diff = ImageChops.difference(rgb, Image.new("RGB", rgb.size, bg)).convert("L")
        box = diff.point(lambda v: 255 if v > tolerance else 0).getbbox()
    if not box:
        return None
    l, t, r, b = box
    if (r - l) * (b - t) > 0.97 * w * h:
        return None
    pad = max(2, int(min(w, h) * 0.01))
    return max(0, l - pad), max(0, t - pad), min(w, r + pad), min(h, b + pad)


def _px_box(spec, iw, ih):
    """[l, t, r, b] crop or {x, y, w, h} region in pixels or 0–1 fractions → pixel box."""
    if isinstance(spec, dict):
        vals = [spec.get(k, 0) for k in ("x", "y", "w", "h")]
        frac = all(0 <= float(v) <= 1.0 for v in vals)
        x, y, w, h = (float(v) * (iw if i % 2 == 0 else ih) if frac else float(v) for i, v in enumerate(vals))
        return x, y, x + w, y + h
    l, t, r, b = (float(v) for v in spec)
    if all(0 <= v <= 1.0 for v in (l, t, r, b)):  # fractions trimmed from each side
        return l * iw, t * ih, iw - r * iw, ih - b * ih
    return l, t, r, b


def _round_and_shadow(pic, radius, border=None, shadow=True):
    spPr = pic._element.spPr
    geom = spPr.find(qn("a:prstGeom"))
    if radius and geom is not None:
        geom.set("prst", "roundRect")
        av = geom.find(qn("a:avLst"))
        if av is None:
            av = etree.SubElement(geom, qn("a:avLst"))
        for child in list(av):
            av.remove(child)
        gd = etree.SubElement(av, qn("a:gd"))
        gd.set("name", "adj")
        gd.set("fmla", "val %d" % int(min(50000, radius / float(max(1, min(pic.width, pic.height))) * 100000)))
    if border:
        pic.line.color.rgb = RGBColor.from_string(border)
        pic.line.width = Emu(int(EMU_PT * 0.75))
    if shadow:
        _shadow(spPr)


def _shadow(spPr):
    for old in spPr.findall(qn("a:effectLst")):
        spPr.remove(old)
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(EMU_PT * 9), dist=str(EMU_PT * 2), dir="5400000",
                          algn="t", rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val="000000")
    etree.SubElement(clr, qn("a:alpha"), val="20000")


def _add_picture(shapes, path, x, y, w, h, crop_px, iw, ih):
    pic = shapes.add_picture(str(path), Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    l, t, r, b = crop_px
    pic.crop_left, pic.crop_top = l / float(iw), t / float(ih)
    pic.crop_right, pic.crop_bottom = (iw - r) / float(iw), (ih - b) / float(ih)
    return pic


def draw_image(r, slide, rect, b):
    th = r.theme
    x, y, w, h = rect
    path = b["path"]
    with Image.open(str(path)) as im:
        iw, ih = im.size
    crop = (0.0, 0.0, float(iw), float(ih))
    if b.get("crop"):
        crop = _px_box(b["crop"], iw, ih)
    elif b.get("trim", True) and b.get("style", "card") != "plain":
        tb = trim_box(path)
        if tb:
            crop = tuple(float(v) for v in tb)
    cl, ct, cr, cb = crop
    cl, ct = max(0.0, cl), max(0.0, ct)
    cr, cb = min(float(iw), cr), min(float(ih), cb)
    if cr - cl < 4 or cb - ct < 4:
        r.warnings.append(u"%s：crop 範圍不對，已改用整張圖" % r._where)
        cl, ct, cr, cb = 0.0, 0.0, float(iw), float(ih)
    style = b.get("style") or "card"
    notes = [hl for hl in b.get("highlights") or [] if hl.get("note")]
    # side panel for the numbered notes
    panel_w = 0
    if notes:
        panel_w = int(min(w * 0.34, EMU_IN * 3.4))
        w_img = w - panel_w - int(EMU_PT * 16)
    else:
        w_img = w
    cap_size = th.snap(th.size("caption") + 2)
    cap_h = int(cap_size * EMU_PT * 2.2) if b.get("caption") else 0
    bar_h = int(min(EMU_IN * 0.3, h * 0.08)) if style == "browser" else 0
    box_w, box_h = w_img, h - cap_h - bar_h
    aspect = (cr - cl) / float(cb - ct)
    if b.get("fit") == "cover":
        dw, dh = box_w, box_h
        want = box_w / float(box_h)
        if aspect > want:  # too wide: crop the sides
            extra = (cr - cl) - (cb - ct) * want
            cl, cr = cl + extra / 2.0, cr - extra / 2.0
        else:
            extra = (cb - ct) - (cr - cl) / want
            ct, cb = ct + extra / 2.0, cb - extra / 2.0
    else:
        dw = min(box_w, box_h * aspect)
        dh = dw / aspect
    px = x + (w_img - dw) / 2.0
    total_h = dh + bar_h
    py = y + (h - cap_h - total_h) / 2.0 + bar_h
    radius = EMU_PT * 6

    cv = Canvas(r, slide, "Image")  # frame, picture, annotations and caption move as one object
    if style == "browser":
        frame = cv.rect(px, py - bar_h, dw, dh + bar_h, fill=th.neutral(100), stroke=th.neutral(300), width=0.75,
                        radius=radius)
        _shadow(frame._element.spPr)
        dot = bar_h * 0.13
        for i, c in enumerate(DOTS):
            cv.ellipse(px + bar_h * 0.45 + i * dot * 3.2, py - bar_h / 2.0, dot, dot, fill=c)
        if b.get("url"):
            uw = min(dw * 0.6, dw - bar_h * 3)
            pill = cv.rect(px + (dw - uw) / 2.0, py - bar_h * 0.78, uw, bar_h * 0.56, fill="FFFFFF", radius=bar_h * 0.28)
            cv.text(0, 0, b["url"], max(9.0, min(11.0, bar_h * 0.3 / EMU_PT)), color=th.neutral(500), shape=pill)
    pic = _add_picture(cv.shapes, path, px, py, dw, dh, (cl, ct, cr, cb), iw, ih)
    if style == "card":
        _round_and_shadow(pic, radius, border=th.neutral(300))
    elif style == "shadow":
        _round_and_shadow(pic, 0, shadow=True)
    elif style == "browser":
        _round_and_shadow(pic, 0, border=None, shadow=False)

    def to_slide(bx):
        sx, sy = dw / (cr - cl), dh / (cb - ct)
        l, t, rr, bb = bx
        return px + (l - cl) * sx, py + (t - ct) * sy, (rr - l) * sx, (bb - t) * sy

    # zoom inset
    if b.get("zoom"):
        z = b["zoom"]
        zl, zt, zr, zb = _px_box(z, iw, ih)
        if zr - zl > 2 and zb - zt > 2:
            ox, oy, ow, oh = to_slide((zl, zt, zr, zb))
            cv.rect(ox, oy, ow, oh, stroke=th.primary, width=1.5, dash=True)
            iw2 = dw * float(z.get("size", 0.42))
            ih2 = iw2 * (zb - zt) / float(zr - zl)
            if ih2 > dh * 0.6:
                ih2 = dh * 0.6
                iw2 = ih2 * (zr - zl) / float(zb - zt)
            pos = z.get("position") or ("bottom-left" if ox + ow / 2.0 > px + dw / 2.0 else "bottom-right")
            ix = px + dw - iw2 - EMU_PT * 10 if "right" in pos else px + EMU_PT * 10
            iy = py + dh - ih2 - EMU_PT * 10 if "bottom" in pos else py + EMU_PT * 10
            inset = _add_picture(cv.shapes, path, ix, iy, iw2, ih2, (zl, zt, zr, zb), iw, ih)
            inset.line.color.rgb = RGBColor.from_string(th.primary)
            inset.line.width = Emu(int(EMU_PT * 2))
            _shadow(inset._element.spPr)
            if (zr - zl) / max(0.01, iw2 / float(EMU_IN)) < 96:
                r.warnings.append(u"%s：放大區解析度不足，放大後會模糊（請用更清楚的截圖，或縮小 zoom.size）" % r._where)
    # highlights
    hls = b.get("highlights") or []
    if hls:
        badge = EMU_PT * 9
        num_size = th.snap(11)
        n = 0
        for hl in hls:
            color = resolve_color(hl.get("color") or "alert", th) or th.alert
            hx, hy, hw, hh = to_slide(_px_box(hl, iw, ih))
            cv.rect(hx, hy, hw, hh, stroke=color, width=2.25, radius=EMU_PT * 3)
            label = hl.get("label")
            if label is None and hl.get("note"):
                n += 1
                label = str(n)
            hl["_label"] = label
            if label:
                bx = max(x, hx - badge * 0.9)
                by = max(y, hy - badge * 0.9)
                circle = cv.ellipse(bx + badge, by + badge, badge, badge, fill=color, stroke="FFFFFF", width=1.5)
                cv.text(0, 0, str(label), num_size, color="FFFFFF", bold=True, align="middle", shape=circle)
        if notes:
            nx = x + w - panel_w
            size = th.snap(th.size("body", 2))
            line_h = size * EMU_PT * 1.35
            est = []
            from .renderer import em_width, plain_text
            for hl in notes:
                lines = max(1, int(-(-em_width(plain_text(hl["note"])) * size * EMU_PT // max(1, panel_w - badge * 3))))
                est.append(lines * line_h)
            ny = y + max(0.0, (h - cap_h - (sum(est) + EMU_PT * 12 * (len(notes) - 1))) / 2.0)
            for hl, eh in zip(notes, est):
                color = resolve_color(hl.get("color") or "alert", th) or th.alert
                circle = cv.ellipse(nx + badge, ny + badge, badge, badge, fill=color)
                cv.text(0, 0, str(hl["_label"] or ""), num_size, color="FFFFFF", bold=True, align="middle", shape=circle)
                cv.text(nx + badge * 2.8, ny + badge - line_h / 2.0, hl["note"], size, color=th.neutral(900),
                        width=panel_w - badge * 2.8, height=eh, wrap=True)
                ny += eh + EMU_PT * 12
    if b.get("caption"):
        cv.text(px + dw / 2.0, py + dh + EMU_PT * 8, b["caption"], cap_size, color=th.neutral(500), align="middle",
                width=max(dw, EMU_IN * 2), wrap=True)
    dpi = (cr - cl) / max(0.01, dw / float(EMU_IN))
    if dpi < 96:
        r.warnings.append(u"%s：圖片「%s」解析度偏低（約 %d dpi），投影時會模糊；請用更清楚的截圖，或縮小顯示區"
                          % (r._where, getattr(path, "name", str(path).replace("\\", "/").split("/")[-1]), dpi))


def draw_gallery(r, slide, rect, b):
    items = b["items"]
    n = len(items)
    if not n:
        return
    x, y, w, h = r._chart_title(slide, rect, b["title"]) if b.get("title") else rect
    cols = int(b.get("columns") or (n if n <= 3 else (2 if n == 4 else 3)))
    rows = -(-n // cols)
    gap = r.gap
    cw = (w - gap * (cols - 1)) / float(cols)
    ch = (h - gap * (rows - 1)) / float(rows)
    for i, it in enumerate(items):
        cx = x + (i % cols) * (cw + gap)
        cy = y + (i // cols) * (ch + gap)
        draw_image(r, slide, (int(cx), int(cy), int(cw), int(ch)), it)


def grid_preview(path, max_edge=1280):
    """Image with a labelled 10% grid so the model can read highlight / crop / zoom coordinates (0–1)."""
    from .sketch import _encode, _flatten

    with Image.open(str(path)) as im:
        size = im.size
        base = _flatten(im)
    base.thumbnail((max_edge, max_edge))
    w, h = base.size
    ov = Image.new("RGBA", base.size, (0, 0, 0, 0))
    d = ImageDraw.Draw(ov)
    for i in range(1, 10):
        x, y = int(w * i / 10.0), int(h * i / 10.0)
        a = 150 if i == 5 else 90
        d.line([(x, 0), (x, h)], fill=(255, 0, 80, a), width=1)
        d.line([(0, y), (w, y)], fill=(255, 0, 80, a), width=1)
        for label, pos in (("%.1f" % (i / 10.0), (x + 2, 2)), ("%.1f" % (i / 10.0), (2, y + 1))):
            d.rectangle([pos[0] - 1, pos[1] - 1, pos[0] + 19, pos[1] + 11], fill=(255, 255, 255, 200))
            d.text(pos, label, fill=(220, 0, 60, 255))
    out = Image.alpha_composite(base.convert("RGBA"), ov).convert("RGB")
    return {"size": size, "data": _encode(out, max_edge)}


def draw_placeholder(r, slide, rect, b):
    """Tidy "material still needed" frame, so a half-collected deck still reads well in review."""
    th = r.theme
    x, y, w, h = rect
    cv = Canvas(r, slide, "Placeholder")
    cv.rect(x, y, w, h, fill=th.neutral(50), stroke=th.neutral(300), width=1.25, dash=True, radius=EMU_PT * 6)
    size = th.snap(th.size("body", 2))
    cv.text(x + w / 2.0, y + h / 2.0 - size * EMU_PT * 0.9, u"待補圖片", th.snap(size + 2), color=th.neutral(500),
            bold=True, align="middle", valign="middle")
    cv.text(x + w / 2.0, y + h / 2.0 + size * EMU_PT * 0.6, b.get("need") or "", size, color=th.neutral(500),
            align="middle", valign="top", width=w * 0.8, wrap=True, height=size * EMU_PT * 3)
