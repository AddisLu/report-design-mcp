"""Visual lint of a built .pptx: geometry, overflow, type-scale consistency, contrast, density, empty space.

Pure python-pptx, so it runs everywhere (the PowerPoint preview images are an extra on Windows).
"""
from __future__ import annotations

from collections import Counter

from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn

from .renderer import EMU_IN, EMU_PT, _master_sizes, _text_height, plain_text
from .template import open_presentation
from .template_inspector import TITLE_TYPES
from .theme import contrast, load_theme

CHROME_PH = (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE)
MAX_CHARS = 320
MAX_SIZES = 10


def _ph_type(shp):
    try:
        return shp.placeholder_format.type if shp.is_placeholder else None
    except Exception:
        return None


def _box(shp):
    if shp.left is None or shp.width is None:
        return None
    return int(shp.left), int(shp.top), int(shp.width), int(shp.height)


def _inter(a, b):
    w = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
    h = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
    return w * h if w > 0 and h > 0 else 0


def _solid_rgb(el):
    """Hex of an explicit <a:solidFill><a:srgbClr> directly under el (spPr / tcPr / bgPr)."""
    if el is None:
        return None
    sf = el.find(qn("a:solidFill"))
    if sf is None:
        return None
    c = sf.find(qn("a:srgbClr"))
    return c.get("val").upper() if c is not None else None


def _layout_bg(slide):
    for part in (slide, slide.slide_layout):
        bg = part._element.find("%s/%s/%s" % (qn("p:cSld"), qn("p:bg"), qn("p:bgPr")))
        if bg is not None:
            return _solid_rgb(bg)
    return "FFFFFF"


def _runs(tf):
    for p in tf.paragraphs:
        for r in p.runs:
            yield p, r


def _run_rgb(r):
    rpr = r._r.find(qn("a:rPr"))
    return _solid_rgb(rpr)


class _Ctx:
    def __init__(self, prs, theme):
        self.prs = prs
        self.theme = theme
        self.sw, self.sh = int(prs.slide_width), int(prs.slide_height)
        self.title_pt, self.body_pts = _master_sizes(prs.slide_masters[0])
        self.sizes = Counter()
        self.issues = []

    def add(self, n, level, msg):
        self.issues.append({"slide": n, "level": level, "msg": msg})


def _default_size(ctx, shp):
    t = _ph_type(shp)
    if t in TITLE_TYPES:
        return ctx.title_pt
    if t is not None:
        return ctx.body_pts[0]
    return 18.0


def _text_items(ctx, shp):
    items = []
    for p in shp.text_frame.paragraphs:
        sizes = [r.font.size.pt for r in p.runs if r.font.size is not None]
        size = max(sizes) if sizes else _default_size(ctx, shp)
        items.append((plain_text("".join(r.text for r in p.runs)), size, 20.0 * p.level))
    return items


def _check_slide(ctx, n, slide):
    sw, sh = ctx.sw, ctx.sh
    bg = _layout_bg(slide)
    is_divider = slide.slide_layout._element.get("showMasterSp") == "0"
    surfaces, content = [], []
    chars = 0
    for shp in slide.shapes:
        box = _box(shp)
        if box is None:
            continue
        t = _ph_type(shp)
        if t in CHROME_PH:
            continue
        x, y, w, h = box
        tol = EMU_PT * 2
        if x < -tol or y < -tol or x + w > sw + tol or y + h > sh + tol:
            ctx.add(n, "warn", u"「%s」超出投影片邊界" % shp.name)
        if shp.shape_type == MSO_SHAPE_TYPE.LINE:
            continue
        if shp.has_text_frame:
            text = shp.text_frame.text.strip()
            fill = _solid_rgb(shp._element.find(qn("p:spPr")))
            if not text:
                if fill:
                    surfaces.append((box, fill))
                continue
            chrome = y > sh - EMU_IN * 0.55 and h < EMU_IN * 0.45
            if chrome:
                continue
            chars += len(plain_text(text))
            items = _text_items(ctx, shp)
            tf = shp.text_frame
            ml = (tf.margin_left or 0) + (tf.margin_right or 0)
            mt = (tf.margin_top or 0) + (tf.margin_bottom or 0)
            est = _text_height(items, max(1.0, (w - ml) / float(EMU_PT))) * EMU_PT
            if est > (h - mt) * 1.12 + EMU_PT * 4:
                ctx.add(n, "warn", u"「%s」文字可能超出框（估計需要 %.1f 吋，框高 %.1f 吋）"
                        % (text[:12], est / EMU_IN, h / float(EMU_IN)))
            content.append((box, min(h, est + mt), shp))
            for p, r in _runs(tf):
                if r.font.size is not None:
                    ctx.sizes[r.font.size.pt] += len(r.text)
                    if r.font.size.pt < 9:
                        ctx.add(n, "warn", u"「%s」字級 %.0fpt 太小" % (r.text[:10], r.font.size.pt))
                rgb = _run_rgb(r)
                if not rgb or not r.text.strip():
                    continue
                under = fill
                if under is None:
                    cx, cy = x + w // 2, y + h // 2
                    for (sx, sy, sw_, sh_), sf in surfaces:
                        if sx <= cx <= sx + sw_ and sy <= cy <= sy + sh_:
                            under = sf
                if under is None:
                    under = bg or "FFFFFF"
                big = (r.font.size.pt if r.font.size else _default_size(ctx, shp)) >= 18
                need = 3.0 if big else 4.5
                ratio = contrast(rgb, under)
                if ratio < need:
                    ctx.add(n, "warn", u"「%s」文字與底色對比 %.1f:1，低於 %.1f:1" % (r.text[:10], ratio, need))
        elif shp.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.CHART, MSO_SHAPE_TYPE.GROUP) \
                or getattr(shp, "has_chart", False) or getattr(shp, "has_table", False):
            content.append((box, h, shp))
            if getattr(shp, "has_table", False):
                for row in shp.table.rows:
                    for cell in row.cells:
                        chars += len(cell.text)
                        cfill = _solid_rgb(cell._tc.tcPr) if cell._tc.tcPr is not None else None
                        for p in cell.text_frame.paragraphs:
                            for r in p.runs:
                                if r.font.size is not None:
                                    ctx.sizes[r.font.size.pt] += len(r.text)
                                rgb = _run_rgb(r)
                                if rgb and cfill and r.text.strip() and contrast(rgb, cfill) < 4.5:
                                    ctx.add(n, "warn", u"表格「%s」文字與底色對比 %.1f:1 不足" % (r.text[:8], contrast(rgb, cfill)))
                                    break
        elif shp.has_text_frame is False and _solid_rgb(shp._element.find(qn("p:spPr"))):
            surfaces.append((box, _solid_rgb(shp._element.find(qn("p:spPr")))))
    # overlaps between content shapes (cards/backgrounds excluded)
    for i in range(len(content)):
        for j in range(i + 1, len(content)):
            a, b = content[i], content[j]
            ra = (a[0][0], a[0][1], a[0][2], a[1])
            rb = (b[0][0], b[0][1], b[0][2], b[1])
            inter = _inter(ra, rb)
            small = min(ra[2] * ra[3], rb[2] * rb[3]) or 1
            if inter > small * 0.08:
                ctx.add(n, "warn", u"「%s」和「%s」重疊" % (a[2].name, b[2].name))
    if chars > MAX_CHARS:
        ctx.add(n, "warn", u"文字 %d 字，太密（建議 %d 字內，或拆頁／改用圖表）" % (chars, MAX_CHARS))
    if not is_divider and content:
        top = min(c[0][1] for c in content if _ph_type(c[2]) not in TITLE_TYPES) if any(
            _ph_type(c[2]) not in TITLE_TYPES for c in content) else None
        body = [c for c in content if _ph_type(c[2]) not in TITLE_TYPES]
        area = sw * sh * 0.62
        used = sum(c[0][2] * c[1] for c in body)
        if top is not None and used < area * 0.3:
            ctx.add(n, "info", u"版面偏空（內容約占 %d%%）：可加入 KPI／圖表，或和相鄰頁合併" % int(used * 100 / area))


def lint_deck(path, cfg=None) -> dict:
    prs = open_presentation(path)
    ctx = _Ctx(prs, load_theme(cfg))
    for n, slide in enumerate(prs.slides, 1):
        _check_slide(ctx, n, slide)
    steps = set(ctx.theme.steps)
    off = sorted(s for s in ctx.sizes if s not in steps)
    if off:
        ctx.add(0, "info", u"字級 %s 不在 theme.yaml 的字級表裡" % u"、".join("%gpt" % s for s in off))
    if len(ctx.sizes) > MAX_SIZES:
        ctx.add(0, "info", u"整份用了 %d 種字級（%s），建議 %d 種以內"
                % (len(ctx.sizes), u"、".join("%g" % s for s in sorted(ctx.sizes)), MAX_SIZES))
    return {"slide_count": len(prs.slides), "issues": ctx.issues}


def format_lint(res) -> str:
    issues = res["issues"]
    warns = [i for i in issues if i["level"] == "warn"]
    if not issues:
        return u"版面檢查：%d 頁，沒有發現問題。" % res["slide_count"]
    lines = [u"版面檢查：%d 頁，%d 個需要修正、%d 個建議" % (res["slide_count"], len(warns), len(issues) - len(warns))]
    for i in sorted(issues, key=lambda x: (x["slide"], x["level"] != "warn")):
        where = u"整份" if i["slide"] == 0 else u"投影片 %d" % i["slide"]
        lines.append(u"- %s %s：%s" % (u"⚠" if i["level"] == "warn" else u"💡", where, i["msg"]))
    return "\n".join(lines)
