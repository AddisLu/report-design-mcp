"""End-to-end smoke test. Run from the repo root:  python tests/smoke_test.py"""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

TMP = tempfile.mkdtemp(prefix="report_design_test_")
DESIGN = os.path.join(TMP, "design_system")
shutil.copytree(os.path.join(ROOT, "design_system"), DESIGN)
os.environ["REPORT_DESIGN_DIR"] = DESIGN  # keep the real design_system untouched
os.environ["PYTHONIOENCODING"] = "utf-8"
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from pptx import Presentation  # noqa: E402

from report_engine import deck, tools  # noqa: E402
from report_engine.facts import Resolver  # noqa: E402
from report_engine.sketch import load_sketch_pages, make_canvas  # noqa: E402
from report_engine.sources import load_sources  # noqa: E402

CSV = os.path.join(ROOT, "design_system", "examples", "sample_projects.csv")
EXAMPLE = os.path.join(ROOT, "design_system", "examples", "executive_example.json")
MANAGER = os.path.join(ROOT, "design_system", "examples", "manager_example.json")
BENCH = os.path.join(ROOT, "benchmark", "mcp_reference_spec.json")


def check(cond, msg):
    if not cond:
        raise AssertionError(msg)
    print("ok -", msg)


def deck_texts(path):
    out = []
    for slide in Presentation(path).slides:
        parts = []
        for shp in slide.shapes:
            if shp.has_text_frame:
                parts.append(shp.text_frame.text)
            if getattr(shp, "has_table", False) and shp.has_table:
                parts.extend(c.text for r in shp.table.rows for c in r.cells)
        out.append("\n".join(parts))
    return out


def test_facts():
    b = load_sources([CSV])
    t = b.tables["t1"]
    check(t.kinds["目標進度"] == "percent" and t.kinds["預算"] == "number", "CSV column kinds detected")
    r = Resolver(b.tables, {"gap": "=t1[代號=B].目標進度 - t1[代號=B].實際進度"})
    errs = []
    check(r.render("{t1[代號=B].實際進度}", errs) == "61%", "single cell percent")
    check(r.render("{=gap*100 | .0f}", errs) == "14", "named fact + expression + format")
    check(r.render("{t1.預算.sum}", errs) == "4,300", "sum with thousands separator")
    check(r.render("{t1[狀態!=正常].count}", errs) == "2", "!= filter + count")
    check(r.render("{t1[實際進度<70].count}", errs) == "3", "percent comparison without % sign")
    check(not errs, "no errors so far")
    r.render("{t1[代號=Z].實際進度}", errs)
    r.render("{t1.不存在}", errs)
    check(len(errs) == 2, "bad references are reported, not invented")


def test_build_and_patch():
    with open(EXAMPLE, encoding="utf-8") as f:
        spec = json.load(f)
    res = deck.build(spec, os.path.join(TMP, "out", "exec.pptx"))
    check(res["ok"], "example spec builds: %s" % res.get("errors"))
    texts = deck_texts(res["path"])
    check(len(texts) == 5, "cover + custom + 2 repeat pages + decision = 5 slides (got %d)" % len(texts))
    joined = "\n".join(texts)
    check("{" not in joined and "}" not in joined, "every {reference} resolved")
    check("供應鏈整合：延遲" in joined and "產線檢測：注意" in joined, "repeat pages expanded per project")
    check("14 個百分點" in joined, "computed gap rendered")
    check(len(res["annotations"]) == 2, "annotation table carried through")
    res2 = deck.patch(res["spec_path"], {"4": {"title": "需要決策：B 案外包（修改後）"}})
    check(res2["ok"], "patch_deck rebuilds")
    texts2 = deck_texts(res2["path"])
    check(len(texts2) == 5 and "修改後" in texts2[4] and texts2[1] == texts[1], "patch changed only the target slide")


def test_overflow_and_layouts():
    bullets = ["第 %d 項：這是一段很長的說明文字，用來測試自動縮字與自動拆成續頁的功能是否正常運作" % i for i in range(1, 26)]
    spec = {"title": "overflow", "audience": "engineer", "sources": [CSV], "slides": [
        {"title": "很多條列", "bullets": bullets},
        {"title": "三個區塊", "bullets": ["a", "b"], "kpis": [{"value": "{t1.count}", "label": "專案數"}],
         "chart": {"type": "column", "from": "t1", "category": "專案", "values": ["預算", "已用"], "sort": "-預算"}},
    ]}
    res = deck.build(spec, os.path.join(TMP, "out"))
    check(res["ok"], "overflow spec builds: %s" % res.get("errors"))
    check(len(res["mapping"][0]["slides"]) >= 2, "long text split into continuation slides")
    check(any("自訂排版" in w for w in res["warnings"]), "too many blocks -> auto custom layout")


def test_errors():
    bad = {"audience": "主管", "sources": [CSV], "slides": [
        {"title": "{t1[代號=Z].專案}"},
        {"layout": "custom", "regions": [{"type": "kpi", "content": [{"value": "1"}]}]},
        {"chart": {"type": "line", "from": "t1", "category": "專案", "values": ["不存在"]}},
    ]}
    res = deck.build(bad, os.path.join(TMP, "out"))
    check(not res["ok"] and len(res["errors"]) == 3, "unfixable problems are all returned at once (%d)" % len(res["errors"]))


def test_sketch_and_blueprint():
    paths = make_canvas(os.path.join(TMP, "canvas.pdf"), pages=3, requirements_page=True)
    check(len(load_sketch_pages(paths[0])) == 4, "multi-page PDF canvas readable")
    check([p["page"] for p in load_sketch_pages(paths[0], "2-3")] == [2, 3], "page ranges")
    r = tools.h_save_layout({"name": "月報-主管版", "slides": [{"regions": [
        {"id": "k", "type": "KPI", "x": 0, "y": 1.5, "w": 12, "h": 2},
        {"id": "b", "type": "重點", "x": 0, "y": 4, "w": 12, "h": 4}]}]})
    check(not r["isError"], "save_layout")
    spec = {"title": "bp", "audience": "manager", "sources": [CSV], "slides": [
        {"layout": "月報-主管版/1", "title": "用已存版型",
         "fill": {"k": [{"value": "{t1.count}", "label": "專案數"}], "b": ["第一點", "第二點"]}}]}
    res = deck.build(spec, os.path.join(TMP, "out"))
    check(res["ok"] and "第一點" in deck_texts(res["path"])[0], "blueprint layout + fill renders")


def _run_sizes(path):
    sizes = set()
    for slide in Presentation(path).slides:
        for shp in slide.shapes:
            frames = []
            if shp.has_text_frame:
                frames.append(shp.text_frame)
            if getattr(shp, "has_table", False) and shp.has_table:
                frames.extend(c.text_frame for r in shp.table.rows for c in r.cells)
            for tf in frames:
                for p in tf.paragraphs:
                    sizes.update(r.font.size.pt for r in p.runs if r.font.size is not None)
    return sizes


def test_design_system():
    from pptx.enum.shapes import PP_PLACEHOLDER
    from report_engine import preview
    from report_engine.lint import lint_deck
    from report_engine.template import resolve_template_path
    from report_engine.theme import Theme, contrast, load_theme, status_key

    th = load_theme()
    check(th.primary == "1E4E8C" and th.font_ea == "Microsoft JhengHei", "theme.yaml loaded")
    check(Theme({"colors": {"primary": "#zzz"}}).primary == "1E4E8C", "invalid hex falls back to default")
    check(th.snap(15.3) == 14 and th.snap(100) == 54, "font sizes snap to the type scale")
    check(contrast(th.status_text("warn"), th.neutral(50)) >= 4.5, "amber status text darkened for contrast")
    check([status_key(x) for x in (u"延遲", "green", u"黃燈", "??")] == ["bad", "good", "warn", None], "status words")

    tpl = resolve_template_path()
    check(tpl.name == "_brand_16x9.pptx", "brand template generated when no company template")
    prs = Presentation(str(tpl))
    check((prs.slide_width, prs.slide_height) == (12192000, 6858000), "brand template is native 16:9")
    names = [L.name for L in prs.slide_masters[0].slide_layouts]
    check(names == [u"封面", u"標題及內容", u"章節", u"兩欄", u"僅標題", u"空白"], "brand layouts: %s" % names)
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    theme_xml = prs.slide_masters[0].part.part_related_by(RT.THEME).blob.decode("utf-8")
    check("Microsoft JhengHei" in theme_xml and "1E4E8C" in theme_xml, "theme XML carries brand font and colour")
    mtime = tpl.stat().st_mtime
    check(resolve_template_path().stat().st_mtime == mtime, "brand template not rebuilt when theme unchanged")
    theme_file = os.path.join(DESIGN, "theme.yaml")
    with open(theme_file, encoding="utf-8") as f:
        original = f.read()
    with open(theme_file, "w", encoding="utf-8") as f:
        f.write(original.replace('primary: "1E4E8C"', 'primary: "0F766E"'))
    rebuilt = Presentation(str(resolve_template_path()))
    check("0F766E" in rebuilt.slide_masters[0].part.part_related_by(RT.THEME).blob.decode("utf-8"),
          "editing theme.yaml rebuilds the brand template")
    with open(theme_file, "w", encoding="utf-8") as f:
        f.write(original)

    for spec_path, n in ((MANAGER, 5), (BENCH, 9)):
        with open(spec_path, encoding="utf-8") as f:
            spec = json.load(f)
        res = deck.build(spec, os.path.join(TMP, "ds_out", os.path.basename(spec_path) + ".pptx"),
                         base_dir=os.path.dirname(spec_path))
        check(res["ok"], "%s builds: %s" % (os.path.basename(spec_path), res.get("errors")))
        check(res["slide_count"] == n, "%s -> %d slides (got %d)" % (os.path.basename(spec_path), n, res["slide_count"]))
        warns = [i for i in res["lint"]["issues"] if i["level"] == "warn"]
        check(not warns, "%s passes the layout lint: %s" % (os.path.basename(spec_path), warns))
        off = _run_sizes(res["path"]) - set(th.steps)
        check(not off, "every font size comes from the type scale (off: %s)" % sorted(off))
    texts = deck_texts(res["path"])
    check(u"延遲" in texts[2] and u"王雅婷" in texts[2], "status block rendered from table")
    check(u"落後專案：原因與對策" in texts[5], "section slide rendered")
    check(u"資料來源：" in texts[1] and texts[1].strip().endswith("2"), "footer source + slide number")
    pres = Presentation(res["path"])
    check(not any(s.is_placeholder and s.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER
                  for s in pres.slides[0].shapes), "no slide number on the cover")

    # lint catches overlap, overflow and poor contrast
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    bad = Presentation(str(tpl))
    sl = bad.slides.add_slide(bad.slide_layouts[5])
    for i in range(2):
        tb = sl.shapes.add_textbox(Inches(1 + i * 0.5), Inches(2), Inches(4), Inches(1))
        r = tb.text_frame.paragraphs[0].add_run()
        r.text = u"重疊的文字方塊" * (12 if i else 1)
        r.font.size = Pt(18)
        r.font.color.rgb = RGBColor.from_string("DDDDDD")
    bad_path = os.path.join(TMP, "bad.pptx")
    bad.save(bad_path)
    msgs = " ".join(i["msg"] for i in lint_deck(bad_path)["issues"])
    check(u"重疊" in msgs and u"超出框" in msgs and u"對比" in msgs, "lint flags overlap / overflow / contrast")
    if os.name != "nt":
        imgs, err = preview.export_images(res["path"], None, 3)
        check(imgs is None and err, "preview degrades gracefully without PowerPoint")


def test_mcp_stdio():
    with open(EXAMPLE, encoding="utf-8") as f:
        spec = json.load(f)
    msgs = [
        {"jsonrpc": "2.0", "id": 1, "method": "initialize", "params": {"protocolVersion": "2025-06-18", "capabilities": {}}},
        {"jsonrpc": "2.0", "method": "notifications/initialized"},
        {"jsonrpc": "2.0", "id": 2, "method": "tools/list"},
        {"jsonrpc": "2.0", "id": 3, "method": "tools/call",
         "params": {"name": "get_design_context", "arguments": {"audience": "主管", "sources": [CSV]}}},
        {"jsonrpc": "2.0", "id": 4, "method": "tools/call",
         "params": {"name": "build_deck", "arguments": {"spec": spec, "out_path": os.path.join(TMP, "mcp")}}},
    ]
    data = "".join(json.dumps(m, ensure_ascii=False) + "\n" for m in msgs).encode("utf-8")
    proc = subprocess.run([sys.executable, os.path.join(ROOT, "mcp_server.py")], input=data,
                          stdout=subprocess.PIPE, stderr=subprocess.PIPE, cwd=ROOT, env=os.environ.copy(), timeout=180)
    replies = [json.loads(line) for line in proc.stdout.decode("utf-8").splitlines() if line.strip()]
    check([r.get("id") for r in replies] == [1, 2, 3, 4], "one reply per request, none for the notification")
    check(replies[0]["result"]["serverInfo"]["name"] == "report-design", "initialize")
    check(len(replies[1]["result"]["tools"]) == 8, "tools/list returns 8 tools")
    ctx = replies[2]["result"]["content"][0]["text"]
    check("t1 ＝ sample_projects.csv" in ctx and "高階主管" in ctx, "context includes data and audience")
    check(u"視覺設計指引" in ctx and "preview_deck" in ctx, "context includes the visual design guide")
    check(not replies[3]["result"]["isError"] and "已產出" in replies[3]["result"]["content"][0]["text"],
          "build_deck over MCP")


if __name__ == "__main__":
    try:
        for fn in (test_facts, test_build_and_patch, test_overflow_and_layouts, test_errors,
                   test_sketch_and_blueprint, test_design_system, test_mcp_stdio):
            print("==", fn.__name__)
            fn()
        print("\nALL PASSED  (outputs in %s)" % TMP)
    except Exception:
        print("\nFAILED  (outputs in %s)" % TMP)
        raise
